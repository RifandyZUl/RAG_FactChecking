"""
Lembar uji coba pedoman 2.0 (docs/rancangan/rancangan_v2.md A6.5): sekitar 10 butir dari set evaluasi eksternal
lama (sudah sekali pakai), dilabeli pemilik proyek dengan pedoman draf sambil mencatat waktu dan bagian yang
membuat ragu. Hasilnya EKSPLORATIF dan tidak dihitung sebagai metrik.

Pemilihan: acak berbenih (SEED) dan terstratifikasi menurut dua sumbu: label final evaluasi eksternal (positif /
negatif) dan apakah setidaknya satu kandidat berlabel PENIPUAN di TurnBackHoax (metadata artikel). Label final
hanya dipakai untuk strata; ia tidak ditulis ke lembar. Kandidat tiap butir = kandidat yang DULU ditampilkan
(P1: konfirmasi_50.csv; P2: peta blok templat 49), teksnya diambil dari articles.json saat ini. Urutan butir dan
kandidat di lembar diacak dengan random.SystemRandom; peta ke butir asli hanya disimpan di kunci.

Yang TIDAK ditulis ke lembar: label lama atau final, keputusan dan catatan lama, keluaran sistem, skor. KEPUTUSAN,
ragu, catatan wajib, menit, dan bagian pedoman yang membuat ragu diisi pemilik proyek sendiri, tanpa bantuan AI
(Aturan Wajib #5). Kode tidak pernah mengisi, mengusulkan, atau mengoreksi sel masukan.

Pemakaian (dari root proyek):
  python -m labeling.sessions.guide_trial --make   # menolak menimpa; menulis kunci dan sidik jari templat terpisah
"""

import argparse
import hashlib
import json
import random
from collections import defaultdict
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from paths import DATA_DIR

TRIAL_DIR = DATA_DIR / "candidates" / "uji_coba_pedoman_v2"
XLSX_PATH = TRIAL_DIR / "uji_coba_pedoman_v2.xlsx"
FILLED_PATH = TRIAL_DIR / "uji_coba_pedoman_v2_terisi.xlsx"
KEY_PATH = TRIAL_DIR / "kunci_uji_coba.json"  # peta nomor -> butir asli + strata; jangan dibuka sebelum selesai
TEMPLATE_PATH = TRIAL_DIR / "sidik_jari_templat_uji_coba.json"
LABEL_FINAL_PATH = DATA_DIR / "candidates" / "sesi_buta_eksternal" / "label_final.json"
SEED = 20261027
# (label final, ada kandidat PENIPUAN) -> jumlah butir; pola penipuan diberi bobot lebih karena aturan barunya di sana
ALLOCATION: dict[tuple[str, str], int] = {("positif", "penipuan"): 3, ("negatif", "penipuan"): 3,
                                          ("positif", "non_penipuan"): 2, ("negatif", "non_penipuan"): 2}
SHEET = "uji_coba"
GENERATED_BY = "labeling.sessions.guide_trial (openpyxl; tata letak dihasilkan kode, sel kuning diisi manusia)"
CAND_HEADER = ["kandidat", "TBH_id", "TBH_tanggal", "TBH_label", "TBH_judul", "TBH_narasi", "TBH_kesimpulan", "TBH_referensi"]
WIDTHS = [22, 16, 13, 14, 36, 80, 60, 40]
NO_TEXT = "(TIDAK ADA DI SUMBER)"
GUIDE = [
    "Berkas ini disiapkan kode (labeling.sessions.guide_trial). Sel kuning diisi pemilik proyek sendiri, tanpa bantuan AI.",
    "Tujuan: uji coba pedoman 2.0 (draf) sebelum dikunci. Hasilnya eksploratif; bukan metrik dan bukan label set uji.",
    "Pedoman: testset/ANNOTATION_GUIDE_v2.md. Tulis commit pedoman yang dipakai di sel kuning paling atas.",
    "KEPUTUSAN: id artikel TurnBackHoax yang klaimnya SAMA (beberapa id dipisah koma bila lebih dari satu artikel sah, aturan 7) | TIDAK ADA.",
    "ragu: YA atau kosong. menit: lama menilai butir itu, dalam menit (perkiraan wajar).",
    "catatan wajib: isi sesuai pedoman 2.0 bagian 6. bagian yang membuat ragu: nomor aturan atau bagian pedoman, dan alasannya.",
    ("Daftar acuan saluran resmi belum ada pada uji coba ini. Untuk saluran resmi, pakai hanya Kesimpulan dan Referensi "
     "artikel; catat di 'bagian yang membuat ragu' bila itu tidak cukup."),
    "Urutan butir dan kandidat diacak; tidak ada skor, label lama, atau keluaran sistem. Butir ini pernah Anda labeli pada evaluasi eksternal; jangan membuka berkas lama.",
    "Lembar terkunci tanpa kata sandi. Jangan menyisipkan atau menghapus baris.",
    f"Simpan hasil sebagai {FILLED_PATH.name} di folder yang sama; berkas ini tidak ditimpa. {KEY_PATH.name}: JANGAN dibuka.",
]
INPUT_LABELS = ("KEPUTUSAN", "ragu (YA / kosong)", "menit", "catatan wajib", "bagian pedoman yang membuat ragu")


def stratum(label_final: dict[str, Any], candidate_labels: list[str]) -> tuple[str, str]:
    """Strata satu butir: (positif|negatif menurut label final, penipuan|non_penipuan menurut label kandidat)."""
    kind = "positif" if label_final["expected_verdict"] == "ditemukan" else "negatif"
    return kind, "penipuan" if "PENIPUAN" in candidate_labels else "non_penipuan"


def select(strata: dict[str, tuple[str, str]], seed: int = SEED,
           allocation: dict[tuple[str, str], int] = ALLOCATION) -> list[str]:
    """Kode butir terpilih: per strata, `random.Random(seed).sample` atas kode yang diurutkan. Deterministik."""
    rng = random.Random(seed)
    groups: dict[tuple[str, str], list[str]] = defaultdict(list)
    for code, s in strata.items():
        groups[s].append(code)
    chosen = []
    for s, k in allocation.items():
        pool = sorted(groups.get(s, []), key=lambda c: (c[:2], int(c.split("-")[1])))
        if len(pool) < k:
            raise SystemExit(f"strata {s} hanya punya {len(pool)} butir, perlu {k}")
        chosen += rng.sample(pool, k)
    return chosen


def load_source() -> tuple[dict[str, dict[str, Any]], dict[str, tuple[str, str]]]:
    """Butir evaluasi eksternal yang tidak dikeluarkan: pesan Liputan6 + kandidat yang dulu tampil, dan strata."""
    from labeling.gap import gap_blocks
    from labeling.gap.gap_review import load_valid_sample
    from labeling.gap.gap_workbook import CSV_PATH, read_csv_rows
    from labeling.sessions.blind_workbook import article_lookup

    _, _, valid = load_valid_sample()
    by_id = article_lookup()
    p1: dict[int, list[str]] = defaultdict(list)
    p1_title: dict[int, str] = {}
    for r in read_csv_rows(CSV_PATH):
        p1[int(r["butir"])].append(r["TBH_id"])
        p1_title.setdefault(int(r["butir"]), r["L6_judul"])
    p2 = {b["butir"]: b for b in json.loads(gap_blocks.TEMPLATE_PATH.read_text(encoding="utf-8"))["blok"]}
    items, strata = {}, {}
    for b in json.loads(LABEL_FINAL_PATH.read_text(encoding="utf-8"))["butir"]:
        if b["dikeluarkan"]:
            continue
        l6 = valid[b["l6_id"]]
        if b["putaran"] == 1:
            if p1_title[b["butir"]].replace("\r", "") != l6["judul"].replace("\r", ""):
                raise SystemExit(f"{b['kode']}: judul Liputan6 tidak sama dengan konfirmasi_50.csv")
            ids = p1[b["butir"]]
        else:
            if p2[b["butir"]]["l6_id"] != b["l6_id"]:
                raise SystemExit(f"{b['kode']}: l6_id tidak sama dengan templat 49")
            ids = [c["id"] for c in p2[b["butir"]]["kandidat"]]
        items[b["kode"]] = {
            "judul": l6["judul"], "tanggal": l6["tanggal"], "pesan": l6["pesan"],
            "kandidat": [{"id": i, "tanggal": by_id[i]["date"], "label": by_id[i]["label"], "judul": by_id[i]["title"],
                          "narasi": by_id[i].get("narasi") or "", "kesimpulan": by_id[i].get("kesimpulan") or "",
                          "referensi": "\n".join(by_id[i].get("references") or [])} for i in ids]}
        strata[b["kode"]] = stratum(b["label_final"], [by_id[i]["label"] for i in ids])
    return items, strata


def layout(items: list[dict[str, Any]]) -> tuple[list[list[str]], list[dict[str, Any]], list[tuple[int, int]]]:
    """Kisi teks (baris x 8 kolom), peta blok (tanpa kode butir asli), dan daftar sel masukan (baris, kolom)."""
    grid: list[list[str]] = []
    inputs: list[tuple[int, int]] = []
    blocks = []

    def row(*cells: str) -> int:
        grid.append([*cells, *[""] * (len(CAND_HEADER) - len(cells))])
        return len(grid)

    inputs.append((row("commit pedoman yang dipakai", ""), 2))
    row()
    for n, item in enumerate(items, 1):
        row(f"BUTIR {n}", f"Liputan6 terbit {item['tanggal']}")
        row("L6_judul", item["judul"])
        row("L6_kutipan_pesan", item["pesan"])
        r = row(INPUT_LABELS[0], "", INPUT_LABELS[1], "", INPUT_LABELS[2], "")
        inputs += [(r, 2), (r, 4), (r, 6)]
        rows = {"keputusan": r}
        for label in INPUT_LABELS[3:]:
            rr = row(label, "")
            inputs.append((rr, 2))
            rows[label] = rr
        row(*CAND_HEADER)
        for i, c in enumerate(item["kandidat"], 1):
            row(str(i), c["id"], c["tanggal"], c["label"], c["judul"], c["narasi"] or NO_TEXT, c["kesimpulan"] or NO_TEXT,
                c["referensi"] or NO_TEXT)
        row()
        blocks.append({"butir": n, "baris": rows, "kandidat": [c["id"] for c in item["kandidat"]]})
    return grid, blocks, inputs


def evidence_hash(grid: list[list[str]], inputs: list[tuple[int, int]]) -> str:
    """Sidik jari semua sel SELAIN sel masukan: sama untuk templat kosong dan berkas terisi."""
    skip = set(inputs)
    fixed = [["" if (r, c) in skip else cell.replace("\r\n", "\n").replace("\r", "\n") for c, cell in enumerate(line, 1)]
             for r, line in enumerate(grid, 1)]
    return hashlib.sha256(json.dumps(fixed, ensure_ascii=False).encode("utf-8")).hexdigest()


def write_workbook(grid: list[list[str]], inputs: list[tuple[int, int]], path: Path) -> None:
    """Lembar uji coba (terkunci tanpa kata sandi; hanya sel masukan terbuka) dan lembar petunjuk. Semua sel teks."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    wb.properties.creator = GENERATED_BY
    wb.properties.title = "Uji coba pedoman 2.0 (eksploratif)"
    wrap = Alignment(wrap_text=True, vertical="top")
    yellow, blue, grey = (PatternFill("solid", fgColor=c) for c in ("FFF2CC", "BDD7EE", "D9D9D9"))
    for j, width in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(j)].width = width
    last = get_column_letter(len(CAND_HEADER))
    open_cells = set(inputs)
    for r, line in enumerate(grid, 1):
        for c, raw in enumerate(line, 1):
            cell = ws.cell(row=r, column=c)
            text = ILLEGAL_CHARACTERS_RE.sub("", raw)
            if text:
                cell.value = text
                cell.data_type = "s"  # teks berawalan "=" tidak boleh menjadi rumus
            cell.alignment, cell.number_format = wrap, "@"
            if (r, c) in open_cells:
                cell.fill, cell.protection = yellow, Protection(locked=False)
        head = line[0]
        if head.startswith("BUTIR ") or line == CAND_HEADER:
            for c in range(1, len(CAND_HEADER) + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill, cell.font = blue if head.startswith("BUTIR ") else grey, Font(bold=True)
        elif head in ("L6_judul", "L6_kutipan_pesan", *INPUT_LABELS[3:], "commit pedoman yang dipakai"):
            ws.merge_cells(f"B{r}:{last}{r}")
            if head == "L6_kutipan_pesan":
                ws.row_dimensions[r].height = min(409, 15 * max(2, len(line[1]) // 200 + line[1].count("\n") + 1))
            elif head in INPUT_LABELS[3:]:
                ws.row_dimensions[r].height = 60
        elif head.isdigit():
            longest = max(len(line[5]) // 75 + line[5].count("\n"), len(line[6]) // 55 + line[6].count("\n"))
            ws.row_dimensions[r].height = min(409, 15 * max(2, longest + 1))
    ws.freeze_panes = "B1"
    ws.protection.sheet, ws.protection.formatColumns, ws.protection.formatRows = True, False, False
    guide = wb.create_sheet("petunjuk")
    guide.column_dimensions["A"].width = 130
    for k, text in enumerate(GUIDE, 1):
        guide.cell(row=k, column=1, value=text).alignment = wrap
    wb.save(path)


def build(items: dict[str, dict[str, Any]], chosen: list[str], rng: random.Random,
          out_dir: Path = TRIAL_DIR) -> dict[str, Any]:
    """Acak urutan, tulis lembar, kunci, dan sidik jari templat ke `out_dir`. `rng` = SystemRandom pada pemakaian nyata."""
    order = list(chosen)
    rng.shuffle(order)
    shown = []
    for code in order:
        cands = list(items[code]["kandidat"])
        rng.shuffle(cands)
        shown.append({**items[code], "kandidat": cands})
    grid, blocks, inputs = layout(shown)
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx, key_path, fp_path = out_dir / XLSX_PATH.name, out_dir / KEY_PATH.name, out_dir / TEMPLATE_PATH.name
    write_workbook(grid, inputs, xlsx)
    now = datetime.now(UTC).isoformat(timespec="seconds")
    key = {"dibuat_utc": now, "catatan": "peta nomor butir uji coba -> butir asli; JANGAN dibuka sebelum uji coba selesai",
           "peta": [{"butir_uji_coba": b["butir"], "kode": code} for b, code in zip(blocks, order)]}
    key_path.write_text(json.dumps(key, ensure_ascii=False, indent=1), encoding="utf-8")
    fp = {"dibuat_utc": now, "benih_pemilihan": SEED, "alokasi": {f"{a}/{b}": k for (a, b), k in ALLOCATION.items()},
          "butir": len(blocks), "baris_kandidat": sum(len(b["kandidat"]) for b in blocks),
          "sidik_jari_bukti": evidence_hash(grid, inputs),
          "sha256_templat_xlsx": hashlib.sha256(xlsx.read_bytes()).hexdigest(),
          "sha256_kunci": hashlib.sha256(key_path.read_bytes()).hexdigest(),
          "sel_masukan": [list(x) for x in inputs], "blok": blocks}
    fp_path.write_text(json.dumps(fp, ensure_ascii=False, indent=1), encoding="utf-8")
    return fp


def make() -> int:
    if XLSX_PATH.exists() or KEY_PATH.exists() or TEMPLATE_PATH.exists():
        raise SystemExit(f"{XLSX_PATH.name}, {KEY_PATH.name}, atau {TEMPLATE_PATH.name} sudah ada: tidak ditimpa")
    items, strata = load_source()
    chosen = select(strata)
    fp = build(items, chosen, random.SystemRandom())
    print(json.dumps({k: v for k, v in fp.items() if k not in ("blok", "sel_masukan")}, ensure_ascii=False, indent=2))
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--make", action="store_true", required=True, help="buat lembar uji coba (menolak menimpa)")
    ap.parse_args()
    return make()


if __name__ == "__main__":
    raise SystemExit(main())
