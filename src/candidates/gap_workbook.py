"""
Berkas konfirmasi celah "belum ditemukan" dalam bentuk .xlsx (docs/rancangan_v2.md bagian 11), dan
pembacaan keputusan pemilik proyek darinya.

Kode ini hanya mengatur tata letak dan MEMBACA keputusan. Kolom KEPUTUSAN tidak pernah diisi, diusulkan,
atau dikoreksi kode: label final selalu manusia (Aturan Wajib #5). Berkas .xlsx dibuat dari
konfirmasi_50.csv yang sudah ada (butir dan kandidat yang sama; sampel tidak diambil ulang).

Asal-usul (pelajaran kejadian 2026-09-22): berkas buatan openpyxl ber-creator kode. Karena itu sidik jari
templat KOSONG disimpan saat dibuat, dan saat membaca diperiksa bahwa berkas terakhir disimpan aplikasi
lain (Excel) serta bahwa sel selain KEPUTUSAN/catatan tidak berubah dari templat.

Pemakaian (dari root proyek):
  PYTHONPATH=src python -m candidates.gap_workbook --make          # CSV -> .xlsx (menolak menimpa)
  PYTHONPATH=src python -m candidates.gap_workbook --read          # setelah pemilik proyek selesai mengisi
"""

import argparse
import csv
import hashlib
import json
import re
import zipfile
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from candidates.gap_review import (
    COLUMNS,
    DECISION_HELP,
    L6_CACHE,
    OUT_DIR,
    banner_label,
    conclusion_text,
)
from evaluation.metrics import wilson_interval

CSV_PATH = OUT_DIR / "konfirmasi_50.csv"
XLSX_PATH = OUT_DIR / "konfirmasi_50.xlsx"
FINGERPRINT_PATH = OUT_DIR / "sidik_jari_templat_xlsx.json"
KEY_PATH = OUT_DIR / "kunci_kandidat.json"
SHEET = "konfirmasi"
GENERATED_BY = "candidates.gap_workbook (openpyxl; tata letak dihasilkan kode, KEPUTUSAN diisi manusia)"
HUMAN_COLUMNS = ("KEPUTUSAN", "catatan")
WIDTHS = {"butir": 6, "L6_tanggal": 12, "ujung_rentang": 9, "L6_judul": 38, "L6_kutipan_pesan": 60, "KEPUTUSAN": 18,
          "catatan": 26, "kandidat": 9, "TBH_id": 9, "TBH_tanggal": 12, "TBH_label": 13, "TBH_asal": 14, "TBH_judul": 42,
          "TBH_kesimpulan": 70}
MAX_CELL_CHARS = 32767  # batas Excel
THRESHOLD = 0.5  # "lebih dari separuh" (bagian 11.2)
NONE, NOT_HOAX = "TIDAK ADA", "BUKAN PUTUSAN HOAKS"


def content_hash(rows: list[dict[str, str]]) -> str:
    """Sidik jari isi SELAIN kolom yang diisi manusia: sama untuk templat kosong dan berkas yang sudah diisi."""
    fixed = [[str(r[c]) for c in COLUMNS if c not in HUMAN_COLUMNS] for r in rows]
    return hashlib.sha256(json.dumps(fixed, ensure_ascii=False).encode("utf-8")).hexdigest()


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames != COLUMNS:
            raise ValueError(f"kolom {path.name} tidak sesuai: {reader.fieldnames}")
        return list(reader)


def write_workbook(rows: list[dict[str, str]], path: Path) -> dict[str, int]:
    """
    Tulis berkas konfirmasi .xlsx: teks dibungkus, baris judul dibekukan, lebar kolom diatur, urutan kolom
    sama dengan CSV (Kesimpulan paling kanan). Semua sel bertipe teks (teks hoaks berawalan "=" tidak boleh
    menjadi rumus). Menolak bila ada KEPUTUSAN/catatan yang sudah terisi: templat harus keluar kosong.
    """
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    filled = [r["butir"] for r in rows if any(r[c].strip() for c in HUMAN_COLUMNS)]
    if filled:
        raise ValueError(f"KEPUTUSAN/catatan sudah terisi pada butir {sorted(set(filled))}: templat tidak dibuat dari berkas terisi")

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    wb.properties.creator = GENERATED_BY
    wb.properties.title = "Konfirmasi celah 'belum ditemukan' (Liputan6 vs TurnBackHoax)"
    wrap = Alignment(wrap_text=True, vertical="top")
    head_fill = PatternFill("solid", fgColor="D9D9D9")
    decision_fill = PatternFill("solid", fgColor="FFF2CC")
    top = Border(top=Side(style="medium"))

    for j, name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=j, value=name)
        cell.font, cell.fill, cell.alignment = Font(bold=True), head_fill, wrap
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS[name]
    ws.freeze_panes = "A2"

    stats = {"karakter_kendali_dibuang": 0, "sel_dipotong": 0}
    for i, row in enumerate(rows, 2):
        first = bool(row["L6_judul"])
        for j, name in enumerate(COLUMNS, 1):
            text, n = ILLEGAL_CHARACTERS_RE.subn("", row[name])
            stats["karakter_kendali_dibuang"] += n
            if len(text) > MAX_CELL_CHARS:
                text, stats["sel_dipotong"] = text[:MAX_CELL_CHARS], stats["sel_dipotong"] + 1
            cell = ws.cell(row=i, column=j)
            if text:
                cell.value = text
                cell.data_type = "s"
            cell.alignment = wrap
            cell.number_format = "@"  # teks: id yang diketik tidak diubah Excel menjadi angka/tanggal
            if first:
                cell.border = top
                if name == "KEPUTUSAN":
                    cell.fill = decision_fill

    guide = wb.create_sheet("petunjuk")
    guide.column_dimensions["A"].width = 120
    for k, line in enumerate([
        "Berkas ini disiapkan kode (candidates.gap_workbook). KEPUTUSAN diisi pemilik proyek sendiri, tanpa bantuan AI.",
        f"Kolom KEPUTUSAN (sel kuning): {DECISION_HELP}.",
        "Kesamaan klaim menurut testset/ANNOTATION_GUIDE.md v1.0.",
        "Kandidat diurutkan menurut tanggal (terbaru dulu), bukan menurut kemiripan; tidak ada skor.",
        "TBH_asal 'antrean' = artikel yang sudah diambil pembaruan terjadwal tetapi belum di-ingest.",
        "ujung_rentang YA = artikel Liputan6 terbit 14-27 September 2026.",
        "Jangan mengubah kolom lain, menyisipkan/menghapus baris, atau mengurutkan ulang: pembacaan memeriksanya.",
        "kunci_kandidat.json memuat skor: JANGAN dibuka sebelum konfirmasi selesai.",
    ], 1):
        guide.cell(row=k, column=1, value=line).alignment = wrap
    wb.save(path)
    return stats


def read_workbook_rows(path: Path) -> list[dict[str, str]]:
    """Baris lembar konfirmasi sebagai teks; angka bulat (id yang diketik di Excel) menjadi "36729", bukan "36729.0"."""
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True, data_only=True)[SHEET]
    it = ws.iter_rows(values_only=True)
    header = [str(v) if v is not None else "" for v in next(it)][:len(COLUMNS)]
    if header != COLUMNS:
        raise ValueError(f"baris judul berubah: {header}")

    def text(v: Any) -> str:
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v)

    rows = [dict(zip(COLUMNS, (text(v) for v in r[:len(COLUMNS)]))) for r in it]
    return [r for r in rows if any(r.values())]


def provenance(path: Path) -> dict[str, str]:
    """creator / lastModifiedBy / modified dari docProps/core.xml (cara mendeteksi berkas buatan kode, Aturan #5)."""
    with zipfile.ZipFile(path) as z:
        core = z.read("docProps/core.xml").decode("utf-8")
        app = z.read("docProps/app.xml").decode("utf-8") if "docProps/app.xml" in z.namelist() else ""

    def tag(xml: str, name: str) -> str:
        m = re.search(rf"<(?:\w+:)?{name}[^>]*>([^<]*)</(?:\w+:)?{name}>", xml)
        return m.group(1) if m else ""

    return {"creator": tag(core, "creator"), "lastModifiedBy": tag(core, "lastModifiedBy"), "modified": tag(core, "modified"),
            "Application": tag(app, "Application"), "AppVersion": tag(app, "AppVersion")}


def normalize_decision(raw: str) -> str:
    """Rapikan spasi/huruf besar saja; isi keputusan tidak ditafsirkan."""
    return re.sub(r"\s+", " ", raw.strip()).upper().rstrip(".")


def read_decisions(rows: list[dict[str, str]]) -> dict[str, Any]:
    """
    Keputusan per butir dari baris PERTAMA butir. Yang tidak dapat dibaca tanpa menebak (kosong, tulisan lain,
    id di luar kandidat, keputusan di baris bukan-pertama) hanya DILAPORKAN, tidak ditafsirkan.
    """
    items: dict[str, dict[str, Any]] = {}
    for r in rows:
        it = items.setdefault(r["butir"], {"butir": r["butir"], "raw": "", "kandidat": {}, "ujung": False, "lain": []})
        it["kandidat"][r["TBH_id"]] = r["TBH_label"]
        if r["L6_judul"]:
            it["raw"], it["ujung"], it["tanggal"] = r["KEPUTUSAN"], r["ujung_rentang"] == "YA", r["L6_tanggal"]
        elif r["KEPUTUSAN"].strip():
            it["lain"].append((r["TBH_id"], r["KEPUTUSAN"]))
    out, problems = [], []
    for it in items.values():
        d = normalize_decision(it["raw"])
        kind = "kosong" if not d else "tidak_ada" if d == NONE else "bukan_putusan_hoaks" if d == NOT_HOAX else \
            "sama" if d in it["kandidat"] else "id_di_luar_kandidat" if d.isdigit() else "tidak_terbaca"
        if it["lain"]:
            problems.append({"butir": it["butir"], "masalah": "keputusan di baris bukan-pertama", "isi": it["lain"]})
        if kind in ("kosong", "id_di_luar_kandidat", "tidak_terbaca"):
            problems.append({"butir": it["butir"], "masalah": kind, "isi": it["raw"]})
        out.append({"butir": it["butir"], "jenis": kind, "keputusan": d, "ujung_rentang": it["ujung"],
                    "label_tbh": it["kandidat"].get(d, "")})
    return {"butir": out, "perlu_ditanyakan": problems}


def proportion(k: int, n: int) -> dict[str, Any]:
    lo, hi = wilson_interval(k, n)
    status = "tanpa_butir" if n == 0 else "jelas_di_atas_50" if lo > THRESHOLD else "jelas_di_bawah_50" if hi < THRESHOLD else "ambigu"
    return {"tidak_ada": k, "butir_sah": n, "proporsi": round(k / n, 4) if n else None, "wilson95": [round(lo, 4), round(hi, 4)],
            "status": status}


def summarize(decisions: list[dict[str, Any]]) -> dict[str, Any]:
    """Proporsi "tidak ada yang cocok" dari butir sah + interval Wilson 95% dan status menurut aturan bagian 11.2."""
    def part(items: list[dict[str, Any]]) -> dict[str, Any]:
        valid = [d for d in items if d["jenis"] in ("tidak_ada", "sama", "id_di_luar_kandidat")]
        return proportion(sum(d["jenis"] == "tidak_ada" for d in valid), len(valid))

    return {"jenis": dict(Counter(d["jenis"] for d in decisions)), "semua": part(decisions),
            "ujung_rentang": part([d for d in decisions if d["ujung_rentang"]]),
            "di_luar_ujung_rentang": part([d for d in decisions if not d["ujung_rentang"]]),
            "label_tbh_pada_pasangan_sama": dict(Counter(d["label_tbh"] for d in decisions if d["jenis"] == "sama"))}


def make() -> int:
    if XLSX_PATH.exists():
        raise SystemExit(f"{XLSX_PATH} sudah ada: tidak ditimpa (bisa berisi keputusan yang sudah diisi)")
    rows = read_csv_rows(CSV_PATH)
    stats = write_workbook(rows, XLSX_PATH)
    back = read_workbook_rows(XLSX_PATH)
    if content_hash(back) != content_hash(rows) and not any(stats.values()):
        raise SystemExit("isi .xlsx tidak sama dengan CSV setelah dibaca ulang")
    fp = {"dibuat_utc": datetime.now(UTC).isoformat(timespec="seconds"), "sumber": CSV_PATH.name,
          "sha256_csv": hashlib.sha256(CSV_PATH.read_bytes()).hexdigest(),
          "sha256_templat_xlsx": hashlib.sha256(XLSX_PATH.read_bytes()).hexdigest(), "sidik_jari_isi": content_hash(back),
          "baris": len(back), "butir": len({r["butir"] for r in back}),
          "keputusan_terisi_saat_dibuat": sum(bool(r["KEPUTUSAN"].strip()) for r in back),
          "catatan_terisi_saat_dibuat": sum(bool(r["catatan"].strip()) for r in back), "asal_usul": provenance(XLSX_PATH), **stats}
    FINGERPRINT_PATH.write_text(json.dumps(fp, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps(fp, ensure_ascii=False, indent=2))
    return 0


def read(allow_partial: bool) -> int:
    fp = json.loads(FINGERPRINT_PATH.read_text(encoding="utf-8"))
    rows = read_workbook_rows(XLSX_PATH)
    prov = provenance(XLSX_PATH)
    checks = {"isi_selain_keputusan_sama_dengan_templat": content_hash(rows) == fp["sidik_jari_isi"],
              "berkas_berubah_sejak_templat": hashlib.sha256(XLSX_PATH.read_bytes()).hexdigest() != fp["sha256_templat_xlsx"],
              "terakhir_disimpan_aplikasi_lain": bool(prov["lastModifiedBy"])}
    res = read_decisions(rows)
    report: dict[str, Any] = {"asal_usul": prov, "pemeriksaan": checks, "perlu_ditanyakan": res["perlu_ditanyakan"]}
    if not checks["isi_selain_keputusan_sama_dengan_templat"]:
        print(json.dumps(report, ensure_ascii=False, indent=2))
        raise SystemExit("kolom selain KEPUTUSAN/catatan berbeda dari templat: hasil TIDAK dihitung")
    if res["perlu_ditanyakan"] and not allow_partial:
        print(json.dumps(report, ensure_ascii=False, indent=2))
        raise SystemExit("ada butir yang belum/tidak terbaca: hasil TIDAK dihitung (lihat perlu_ditanyakan; --partial untuk melihat sementara)")
    report["ringkasan"] = summarize(res["butir"])
    # tabulasi silang label Liputan6 x TurnBackHoax pada pasangan yang dikonfirmasi SAMA (kunci dibuka setelah konfirmasi)
    order = json.loads(KEY_PATH.read_text(encoding="utf-8"))["urutan_butir"]
    cross = Counter()
    for d in res["butir"]:
        if d["jenis"] == "sama" and (L6_CACHE / f"{order[int(d['butir']) - 1]}.html").exists():
            cross[f"{banner_label(conclusion_text(order[int(d['butir']) - 1])) or '(tanpa banner)'} x {d['label_tbh']}"] += 1
    report["label_liputan6_x_turnbackhoax"] = dict(cross)
    report["keputusan"] = res["butir"]
    out = OUT_DIR / "hasil_konfirmasi.json"
    out.write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k != "keputusan"}, ensure_ascii=False, indent=2))
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--make", action="store_true", help="buat konfirmasi_50.xlsx dari konfirmasi_50.csv")
    g.add_argument("--read", action="store_true", help="baca keputusan pemilik proyek dan hitung hasil")
    ap.add_argument("--partial", action="store_true", help="hitung sementara walau ada butir yang belum terbaca")
    args = ap.parse_args()
    return make() if args.make else read(args.partial)


if __name__ == "__main__":
    raise SystemExit(main())
