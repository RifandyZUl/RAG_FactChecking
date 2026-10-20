"""
Berkas konfirmasi butir SISA celah "belum ditemukan" (docs/rancangan_v2.md bagian 11.6), tata letak blok per
butir, dan pembacaan keputusan pemilik proyek darinya.

Kode ini hanya menyiapkan kandidat/tata letak dan MEMBACA. KEPUTUSAN, ragu, dan catatan tidak pernah diisi,
diusulkan, atau dikoreksi kode (Aturan Wajib #5). Butir = artikel sampel yang sah dan BELUM dikonfirmasi pada
50 butir pertama; kandidat dicari dengan cara yang sama (candidates.gap_review.build_items), diurutkan
menurut tanggal, tanpa skor. Bedanya dari berkas pertama (aturan 11.6): Narasi TurnBackHoax ditampilkan utuh,
ada sel `ragu`, masukan hanya di sel blok, dan lembar Rekap berisi rumus saja serta diproteksi.

Pemakaian (dari root proyek):
  PYTHONPATH=src python -m candidates.gap_blocks --make    # menolak menimpa; menulis kunci skor terpisah
  PYTHONPATH=src python -m candidates.gap_blocks --read    # setelah pemilik proyek selesai; TIDAK membuka kunci
  PYTHONPATH=src python -m candidates.gap_blocks --kunci   # baru setelah --read tersimpan; tidak mengubah label
"""

import argparse
import hashlib
import json
import math
import statistics
from collections import Counter
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from candidates.gap_result import RESULT_PATH, compare_with_key, timing
from candidates.gap_review import (
    OUT_DIR,
    banner_label,
    build_items,
    conclusion_text,
    load_valid_sample,
    tbh_date_key,
)
from candidates.gap_workbook import (
    MAX_CELL_CHARS,
    NONE,
    NOT_HOAX,
    normalize_decision,
    proportion,
    provenance,
    read_decisions,
)

XLSX_PATH = OUT_DIR / "konfirmasi_49.xlsx"
TEMPLATE_PATH = OUT_DIR / "sidik_jari_templat_49.json"
FILLED_PATH = OUT_DIR / "konfirmasi_49_terisi.xlsx"  # salinan yang diisi pemilik proyek; inilah yang dibaca --read
KEY_PATH = OUT_DIR / "kunci_kandidat_49.json"  # skor; hanya dibuka --kunci, setelah hasil --read tersimpan
BLOCK_RESULT_PATH = OUT_DIR / "hasil_konfirmasi_49.json"
KEY_RESULT_PATH = OUT_DIR / "perbandingan_kunci_49.json"
SHEET, SUMMARY = "konfirmasi", "Rekap"
GENERATED_BY = "candidates.gap_blocks (openpyxl; tata letak dihasilkan kode, KEPUTUSAN/ragu/catatan diisi manusia)"
CAND_HEADER = ["kandidat", "TBH_id", "TBH_tanggal", "TBH_label", "TBH_asal", "TBH_judul", "TBH_narasi", "TBH_kesimpulan"]
WIDTHS = [14, 20, 16, 13, 16, 40, 80, 60]
INPUT_COLS = {"KEPUTUSAN": 2, "ragu": 4, "catatan": 6}  # kolom sel masukan pada baris masukan tiap blok
NO_NARASI = "(TIDAK ADA NARASI DI SUMBER)"
DOUBT = "YA"
LINE_POINTS, MAX_ROW_POINTS = 15, 409  # tinggi satu baris teks; batas tinggi baris Excel
FAR_DAYS = 30


def wrapped_height(line: list[str]) -> int:
    """
    Tinggi baris (poin) agar sel terpanjang terbaca utuh saat teks dibungkus: perkiraan dari lebar kolom (satu
    karakter per satuan lebar). Dibatasi tinggi baris maksimum Excel, jadi Narasi yang sangat panjang tetap
    terpotong di layar (isi sel utuh).
    """
    lines = max((sum(max(1, math.ceil(len(par) / width)) for par in cell.split("\n")) for cell, width in zip(line, WIDTHS) if cell),
                default=1)
    return min(MAX_ROW_POINTS, LINE_POINTS * lines)


def layout(items: list[dict[str, Any]], sort_by_date: bool = True) -> tuple[list[list[str]], list[dict[str, Any]]]:
    """
    Isi lembar konfirmasi sebagai kisi teks (baris x 8 kolom) dan peta blok. Blok: judul butir, judul
    Liputan6, kutipan pesan, baris masukan, judul tabel kandidat, kandidat (tanggal terbaru dulu; urutan
    masukan apa adanya bila `sort_by_date` False), baris kosong. Peta blok tidak memuat teks artikel (hanya
    id, tanggal, label, asal, dan nomor baris).
    """
    grid: list[list[str]] = []
    blocks = []

    def row(*cells: str) -> int:
        grid.append([*cells, *[""] * (len(CAND_HEADER) - len(cells))])
        return len(grid)

    for n, item in enumerate(items, 1):
        cands = sorted(item["kandidat"], key=lambda c: (tbh_date_key(c["tanggal"]), c["id"]), reverse=True)             if sort_by_date else list(item["kandidat"])
        row(f"BUTIR {n}", f"Liputan6 terbit {item['tanggal']}", "ujung rentang: YA" if item["ujung_rentang"] else "")
        row("L6_judul", item["judul"])
        row("L6_kutipan_pesan", item["pesan"])
        input_row = row("KEPUTUSAN", "", "ragu (YA / kosong)", "", "catatan", "")
        row(*CAND_HEADER)
        for i, c in enumerate(cands, 1):
            row(str(i), c["id"], c["tanggal"], c["label"], c["asal"], c["judul"], c["narasi"] or NO_NARASI, c["kesimpulan"])
        row()
        blocks.append({"butir": n, "l6_id": item["id"], "l6_tanggal": item["tanggal"], "ujung_rentang": item["ujung_rentang"],
                       "baris_masukan": input_row,
                       "kandidat": [{"id": c["id"], "tanggal": c["tanggal"], "label": c["label"], "asal": c["asal"],
                                     "tanpa_narasi": not c["narasi"]} for c in cands]})
    return grid, blocks


def evidence_hash(grid: list[list[str]], blocks: list[dict[str, Any]]) -> str:
    """Sidik jari seluruh sel lembar konfirmasi SELAIN sel masukan: sama untuk templat kosong dan berkas terisi."""
    inputs = {(b["baris_masukan"], col) for b in blocks for col in INPUT_COLS.values()}
    # akhir baris diseragamkan: Excel boleh menulis ulang CRLF/CR saat menyimpan tanpa mengubah isi
    fixed = [["" if (r, c) in inputs else cell.replace("\r\n", "\n").replace("\r", "\n") for c, cell in enumerate(line, 1)]
             for r, line in enumerate(grid, 1)]
    while fixed and not any(fixed[-1]):
        fixed.pop()
    return hashlib.sha256(json.dumps(fixed, ensure_ascii=False).encode("utf-8")).hexdigest()


GUIDE = [
    "Berkas ini disiapkan kode (candidates.gap_blocks). KEPUTUSAN, ragu, dan catatan diisi pemilik proyek sendiri.",
    "Isi HANYA sel kuning di tiap blok. KEPUTUSAN: id artikel TurnBackHoax yang klaimnya SAMA | TIDAK ADA | BUKAN PUTUSAN HOAKS.",
    "ragu: YA atau kosong; hanya boleh ditandai setelah Narasi kandidat dibaca (aturan B, docs/rancangan_v2.md bagian 11.6).",
    "Aturan A dan B di bagian 11.6 berlaku untuk berkas ini; 50 butir pertama tidak dilabel ulang.",
    "Kandidat diurutkan menurut tanggal (terbaru dulu), bukan menurut kemiripan; tidak ada skor.",
    f"Kandidat tanpa Narasi di sumber ditandai {NO_NARASI}.",
    "Lembar terkunci tanpa kata sandi. Jangan membuka kuncinya, menyisipkan/menghapus baris, atau mengurutkan ulang: pembacaan memeriksanya.",
    "kunci_kandidat_49.json memuat skor: JANGAN dibuka sebelum konfirmasi selesai.",
]


def write_workbook(grid: list[list[str]], blocks: list[dict[str, Any]], path: Path,
                   title: str = "Konfirmasi celah 'belum ditemukan' -- butir sisa (aturan 11.6)",
                   guide_lines: list[str] = GUIDE, generated_by: str = GENERATED_BY) -> dict[str, int]:
    """Tulis lembar konfirmasi (blok) dan lembar Rekap (rumus saja). Kedua lembar diproteksi tanpa kata sandi;
    hanya sel masukan blok yang tidak terkunci. Semua sel bertipe teks."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    wb.properties.creator = generated_by
    wb.properties.title = title
    wrap, bold = Alignment(wrap_text=True, vertical="top"), Font(bold=True)
    fills = {"judul": PatternFill("solid", fgColor="BDD7EE"), "masukan": PatternFill("solid", fgColor="FFF2CC"),
             "kepala": PatternFill("solid", fgColor="D9D9D9")}
    for j, width in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(j)].width = width
    last = get_column_letter(len(CAND_HEADER))
    input_rows = {b["baris_masukan"] for b in blocks}
    stats = {"karakter_kendali_dibuang": 0, "sel_dipotong": 0}
    for r, line in enumerate(grid, 1):
        kind = "judul" if line[0].startswith("BUTIR ") else "kepala" if line == CAND_HEADER else "masukan" if r in input_rows else ""
        for c, raw in enumerate(line, 1):
            text, n = ILLEGAL_CHARACTERS_RE.subn("", raw)
            stats["karakter_kendali_dibuang"] += n
            if len(text) > MAX_CELL_CHARS:
                text, stats["sel_dipotong"] = text[:MAX_CELL_CHARS], stats["sel_dipotong"] + 1
            cell = ws.cell(row=r, column=c)
            if text:
                cell.value = text
                cell.data_type = "s"  # teks hoaks berawalan "=" tidak boleh menjadi rumus
            cell.alignment, cell.number_format = wrap, "@"
            if kind in ("judul", "kepala"):
                cell.font, cell.fill = bold, fills[kind]
            if kind == "masukan":
                cell.font = bold
                if c in INPUT_COLS.values() or c > INPUT_COLS["catatan"]:  # sel catatan digabung sampai kolom terakhir
                    cell.fill, cell.protection, cell.font = fills["masukan"], Protection(locked=False), Font(bold=False)
        if line[0] in ("L6_judul", "L6_kutipan_pesan"):
            ws.merge_cells(f"B{r}:{last}{r}")
            ws.row_dimensions[r].height = 30 if line[0] == "L6_judul" else min(409, 15 * max(2, len(line[1]) // 180 + line[1].count("\n") + 1))
        elif kind == "masukan":
            ws.merge_cells(f"F{r}:{last}{r}")
        elif not kind and line[0].isdigit():  # baris kandidat: tanpa ini Excel menampilkan satu baris dan Narasi terpotong
            ws.row_dimensions[r].height = wrapped_height(line)
    doubt = DataValidation(type="list", formula1=f'"{DOUBT}"', allow_blank=True, showErrorMessage=True,
                           errorTitle="ragu", error="Isi YA atau kosongkan.")
    ws.add_data_validation(doubt)
    for b in blocks:
        doubt.add(ws.cell(row=b["baris_masukan"], column=INPUT_COLS["ragu"]))
    ws.freeze_panes = "B1"
    ws.protection.sheet, ws.protection.formatColumns, ws.protection.formatRows = True, False, False  # lebar/tinggi boleh diubah

    rk = wb.create_sheet(SUMMARY)
    for j, (name, width) in enumerate([("butir", 8), ("KEPUTUSAN", 24), ("ragu", 8), ("", 4), ("ringkasan", 34), ("jumlah", 10)], 1):
        rk.cell(row=1, column=j, value=name).font = bold
        rk.column_dimensions[get_column_letter(j)].width = width
    ref = f"'{SHEET}'!"
    for i, b in enumerate(blocks, 2):
        d, g = f"{ref}{get_column_letter(INPUT_COLS['KEPUTUSAN'])}{b['baris_masukan']}", f"{ref}{get_column_letter(INPUT_COLS['ragu'])}{b['baris_masukan']}"
        rk.cell(row=i, column=1, value=b["butir"])
        rk.cell(row=i, column=2, value=f'=IF({d}="","",{d}&"")')
        rk.cell(row=i, column=3, value=f'=IF({g}="","",{g}&"")')
    rng_d, rng_g, n = f"B2:B{len(blocks) + 1}", f"C2:C{len(blocks) + 1}", len(blocks)
    for i, (name, formula) in enumerate([
        ("butir", f"={n}"), ("sudah diisi", f'=COUNTIF({rng_d},"?*")'), ("belum diisi", f'={n}-COUNTIF({rng_d},"?*")'),
        (NONE, f'=COUNTIF({rng_d},"{NONE}")'), (NOT_HOAX, f'=COUNTIF({rng_d},"{NOT_HOAX}")'),
        ("cocok (id)", f'=COUNTIF({rng_d},"?*")-COUNTIF({rng_d},"{NONE}")-COUNTIF({rng_d},"{NOT_HOAX}")'),
        ("ragu = YA", f'=COUNTIF({rng_g},"{DOUBT}")'),
        ("Lembar ini hanya rumus dan dikunci; isi di lembar konfirmasi.", None),
    ], 2):
        rk.cell(row=i, column=5, value=name)
        if formula:
            rk.cell(row=i, column=6, value=formula)
    rk.freeze_panes = "A2"
    rk.protection.sheet = True  # tanpa kata sandi: mencegah salah ketik, bukan pengaman

    guide = wb.create_sheet("petunjuk")
    guide.column_dimensions["A"].width = 130
    for k, text in enumerate(guide_lines, 1):
        guide.cell(row=k, column=1, value=text).alignment = wrap
    wb.save(path)
    return stats


def read_grid(path: Path) -> list[list[str]]:
    """Lembar konfirmasi sebagai kisi teks; angka bulat (id yang diketik) menjadi "36729", bukan "36729.0"."""
    from openpyxl import load_workbook

    def text(v: Any) -> str:
        if v is None:
            return ""
        return str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)

    ws = load_workbook(path, read_only=True, data_only=True)[SHEET]
    return [[text(v) for v in (*r, *[None] * len(CAND_HEADER))[:len(CAND_HEADER)]] for r in ws.iter_rows(values_only=True)]


def read_inputs(grid: list[list[str]], blocks: list[dict[str, Any]]) -> dict[str, Any]:
    """
    Keputusan dan tanda ragu dari SEL BLOK. Yang dirapikan hanya spasi dan huruf besar-kecil ("Tidak Ada" ->
    "TIDAK ADA"); sel kosong, tulisan lain, id di luar kandidat, dan ragu selain YA/kosong hanya DILAPORKAN.
    """
    def cell(r: int, c: int) -> str:
        return grid[r - 1][c - 1] if r <= len(grid) else ""

    rows, doubt, problems = [], {}, []
    for b in blocks:
        r = b["baris_masukan"]
        raw_doubt = normalize_decision(cell(r, INPUT_COLS["ragu"]))
        if raw_doubt not in ("", DOUBT):
            problems.append({"butir": str(b["butir"]), "masalah": "ragu selain YA/kosong", "isi": cell(r, INPUT_COLS["ragu"])})
        doubt[str(b["butir"])] = raw_doubt == DOUBT
        for i, c in enumerate(b["kandidat"]):  # bentuk baris yang dipahami gap_workbook.read_decisions
            rows.append({"butir": str(b["butir"]), "TBH_id": c["id"], "TBH_label": c["label"], "L6_judul": "x" if i == 0 else "",
                         "KEPUTUSAN": cell(r, INPUT_COLS["KEPUTUSAN"]) if i == 0 else "",
                         "ujung_rentang": "YA" if b["ujung_rentang"] else "", "L6_tanggal": b["l6_tanggal"]})
    res = read_decisions(rows)
    for d in res["butir"]:
        d["ragu"] = doubt[d["butir"]]
        if d["ragu"] and d["jenis"] == "bukan_putusan_hoaks":
            problems.append({"butir": d["butir"], "masalah": "ragu pada BUKAN PUTUSAN HOAKS", "isi": DOUBT})
    return {"butir": res["butir"], "perlu_ditanyakan": res["perlu_ditanyakan"] + problems}


def doubt_bounds(decisions: list[dict[str, Any]]) -> dict[str, Any]:
    """Batas dua arah (11.6): semua butir ragu dihitung TIDAK ADA (atas) atau cocok (bawah). Label utama tetap."""
    valid = [d for d in decisions if d["jenis"] in ("tidak_ada", "sama")]
    k, n = sum(d["jenis"] == "tidak_ada" for d in valid), len(valid)
    up = [int(d["butir"]) for d in valid if d["ragu"] and d["jenis"] == "sama"]
    down = [int(d["butir"]) for d in valid if d["ragu"] and d["jenis"] == "tidak_ada"]
    return {"utama": proportion(k, n), "batas_atas": {**proportion(k + len(up), n), "butir_dibalik": up},
            "batas_bawah": {**proportion(k - len(down), n), "butir_dibalik": down}}


def combine(old: dict[str, Any], new: dict[str, Any]) -> dict[str, Any]:
    """Gabungan 50 butir lama (hasil dan batas 11.5 apa adanya) + butir baru; aturan berhenti 11.6."""
    def add(a: dict[str, Any], b: dict[str, Any]) -> dict[str, Any]:
        return proportion(a["tidak_ada"] + b["tidak_ada"], a["butir_sah"] + b["butir_sah"])

    old_sens = old["sensitivitas_EKSPLORATIF_tidak_dipraregistrasi"]
    main = add(old["hasil_utama_bagian_11_2"], new["utama"])
    verdict = {"jelas_di_atas_50": "layak dikerjakan", "jelas_di_bawah_50": "tidak layak dikerjakan",
               "ambigu": "ambang TIDAK terpenuhi; sampling dihentikan, tidak ada perluasan ketiga"}.get(main["status"], main["status"])
    upper, lower = add(old_sens["batas_atas"], new["batas_atas"]), add(old_sens["batas_bawah"], new["batas_bawah"])
    return {"hasil_utama": main, "keputusan_menurut_11_6": verdict,
            "sensitivitas_EKSPLORATIF": {"batas_atas": upper, "batas_bawah": lower,
                                         "kategori_sama_di_kedua_batas": main["status"] == upper["status"] == lower["status"]}}


def day_gap(l6_iso: str, tbh_ddmmyyyy: str) -> int:
    """Tanggal terbit TurnBackHoax dikurangi tanggal terbit Liputan6, dalam hari (positif = TurnBackHoax belakangan)."""
    return (date(*tbh_date_key(tbh_ddmmyyyy)) - date.fromisoformat(l6_iso)).days


def gap_summary(gaps: list[int]) -> dict[str, Any]:
    """Sebaran selisih hari: bertanda dan mutlak (kuartil metode inklusif), dan berapa yang terpaut jauh."""
    def spread(xs: list[int]) -> dict[str, float] | None:
        if len(xs) < 2:
            return None
        q1, q2, q3 = statistics.quantiles(xs, n=4, method="inclusive")
        return {"min": min(xs), "kuartil_bawah": q1, "median": q2, "kuartil_atas": q3, "maks": max(xs)}

    absolute = [abs(g) for g in gaps]
    return {"pasangan": len(gaps), "bertanda": spread(gaps), "mutlak": spread(absolute),
            "dalam_14_hari": sum(a <= 14 for a in absolute), f"di_atas_{FAR_DAYS}_hari": sum(a > FAR_DAYS for a in absolute),
            "urut": sorted(gaps)}


def crosstab(pairs: list[dict[str, Any]]) -> dict[str, int]:
    return dict(sorted(Counter(f"{p['label_l6']} x {p['label_tbh']}" for p in pairs).items()))


def make() -> int:
    if XLSX_PATH.exists() or KEY_PATH.exists():
        raise SystemExit(f"{XLSX_PATH.name} atau {KEY_PATH.name} sudah ada: tidak ditimpa")
    sample, excluded, valid = load_valid_sample()
    done = list(json.loads(RESULT_PATH.read_text(encoding="utf-8"))["l6_id"].values())
    rest = sorted(set(valid) - set(done))
    if not set(done) <= set(valid) or len(set(done)) != len(done):
        raise SystemExit("butir yang sudah dikonfirmasi tidak cocok dengan sampel sah")
    items, key, n_queue = build_items(rest, valid)
    grid, blocks = layout(items)
    stats = write_workbook(grid, blocks, XLSX_PATH)
    back = read_grid(XLSX_PATH)
    if evidence_hash(back, blocks) != evidence_hash(grid, blocks) and not any(stats.values()):
        raise SystemExit("isi .xlsx tidak sama dengan yang ditulis setelah dibaca ulang")
    KEY_PATH.write_text(json.dumps({"urutan_butir": rest, "skor": key}, ensure_ascii=False, indent=1), encoding="utf-8")
    n_cand = [len(b["kandidat"]) for b in blocks]
    no_narasi = [(b["butir"], c["id"]) for b in blocks for c in b["kandidat"] if c["tanpa_narasi"]]
    fp = {"dibuat_utc": datetime.now(UTC).isoformat(timespec="seconds"), "sampel": len(sample), "dikeluarkan": len(excluded),
          "sah": len(valid), "sudah_dikonfirmasi": len(done), "butir": len(blocks), "baris_kandidat": sum(n_cand),
          "kandidat_per_butir_min_median_maks": [min(n_cand), sorted(n_cand)[len(n_cand) // 2], max(n_cand)],
          "artikel_antrean": n_queue,
          "butir_dengan_kandidat_dari_antrean": sum(any(c["asal"].startswith("antrean") for c in b["kandidat"]) for b in blocks),
          "butir_ujung_rentang": sum(b["ujung_rentang"] for b in blocks),
          "kandidat_tanpa_narasi": len(no_narasi), "kandidat_tanpa_narasi_butir_id": no_narasi,
          "sidik_jari_bukti": evidence_hash(back, blocks), "sha256_templat_xlsx": hashlib.sha256(XLSX_PATH.read_bytes()).hexdigest(),
          "masukan_terisi_saat_dibuat": sum(bool(back[b["baris_masukan"] - 1][c - 1].strip()) for b in blocks for c in INPUT_COLS.values()),
          "asal_usul": provenance(XLSX_PATH), **stats, "blok": blocks}
    TEMPLATE_PATH.write_text(json.dumps(fp, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({k: v for k, v in fp.items() if k != "blok"}, ensure_ascii=False, indent=2))
    return 0


def read(path: Path = FILLED_PATH) -> int:
    fp = json.loads(TEMPLATE_PATH.read_text(encoding="utf-8"))
    blocks, grid, prov = fp["blok"], read_grid(path), provenance(path)
    checks = {"bukti_sama_dengan_templat": evidence_hash(grid, blocks) == fp["sidik_jari_bukti"],
              "berkas_berubah_sejak_templat": hashlib.sha256(path.read_bytes()).hexdigest() != fp["sha256_templat_xlsx"],
              "terakhir_disimpan_aplikasi_lain": bool(prov["lastModifiedBy"])}
    if path != XLSX_PATH and XLSX_PATH.exists():  # templat seharusnya tetap kosong; dilaporkan apa adanya bila tidak
        checks["templat_masih_seperti_saat_dibuat"] = hashlib.sha256(XLSX_PATH.read_bytes()).hexdigest() == fp["sha256_templat_xlsx"]
    res = read_inputs(grid, blocks)
    report: dict[str, Any] = {"berkas": {"terisi": path.name, "sha256_terisi": hashlib.sha256(path.read_bytes()).hexdigest()},
                              "asal_usul": prov, "pemeriksaan": checks, "perlu_ditanyakan": res["perlu_ditanyakan"]}
    if not checks["bukti_sama_dengan_templat"] or res["perlu_ditanyakan"]:
        print(json.dumps(report, ensure_ascii=False, indent=2))
        raise SystemExit("bukti berubah dari templat atau ada butir yang belum/tidak terbaca: hasil TIDAK dihitung")
    decisions, by_n = res["butir"], {str(b["butir"]): b for b in blocks}
    banner = {n: banner_label(conclusion_text(b["l6_id"])) or "(tanpa banner)" for n, b in by_n.items()}
    pairs = []
    for d in decisions:
        if d["jenis"] == "sama":
            c = next(c for c in by_n[d["butir"]]["kandidat"] if c["id"] == d["keputusan"])
            l6_date = by_n[d["butir"]]["l6_tanggal"]
            pairs.append({"butir": int(d["butir"]), "tbh_id": c["id"], "label_l6": banner[d["butir"]], "label_tbh": c["label"],
                          "tbh_asal": c["asal"], "ragu": d["ragu"], "l6_tanggal": l6_date, "tbh_tanggal": c["tanggal"],
                          "tbh_terbit": timing(l6_date, c["tanggal"]), "selisih_hari": day_gap(l6_date, c["tanggal"])})
    old = json.loads(RESULT_PATH.read_text(encoding="utf-8"))
    new = doubt_bounds(decisions)
    combined = combine(old, new)
    statuses = {new[k]["status"] for k in ("batas_atas", "batas_bawah")} | {
        combined["sensitivitas_EKSPLORATIF"][k]["status"] for k in ("batas_atas", "batas_bawah")}
    queue_pairs = [p for p in pairs if p["tbh_asal"].startswith("antrean")]
    report.update({
        "jenis_keputusan": dict(Counter(d["jenis"] for d in decisions)), "butir_ragu": [int(d["butir"]) for d in decisions if d["ragu"]],
        "hasil_sekunder_butir_baru": new,
        "gabungan": combined,
        "kategori_utama_sama_di_semua_batas": statuses == {combined["hasil_utama"]["status"]},
        "putaran_berdampingan_DESKRIPTIF_tanpa_uji": {"lama_50": old["hasil_utama_bagian_11_2"], "baru_49": new["utama"]},
        "bukan_putusan_hoaks": [{"butir": int(d["butir"]), "label_l6": banner[d["butir"]]} for d in decisions
                                if d["jenis"] == "bukan_putusan_hoaks"],
        "ujung_rentang": [{"butir": int(d["butir"]), "jenis": d["jenis"], "keputusan": d["keputusan"]} for d in decisions if d["ujung_rentang"]],
        "label_liputan6_semua_butir": dict(Counter(banner.values())),
        "label_liputan6_pada_tidak_ada": dict(Counter(banner[d["butir"]] for d in decisions if d["jenis"] == "tidak_ada")),
        "label_liputan6_x_turnbackhoax_pasangan_cocok": crosstab(pairs),
        "label_liputan6_x_turnbackhoax_gabungan": crosstab(old["pasangan"] + pairs),
        "tbh_terbit_relatif_liputan6": dict(Counter(p["tbh_terbit"] for p in pairs)),
        "selisih_hari_tbh_dikurangi_liputan6": gap_summary([p["selisih_hari"] for p in pairs]),
        "pasangan_dari_antrean": {"jumlah": len(queue_pairs), "butir_tbh_id": [(p["butir"], p["tbh_id"]) for p in queue_pairs]},
        "tbh_dipilih_lebih_dari_sekali": {i: c for i, c in Counter(p["tbh_id"] for p in pairs).items() if c > 1},
        "tbh_juga_dipilih_pada_putaran_lama": sorted({p["tbh_id"] for p in pairs} & {p["tbh_id"] for p in old["pasangan"]}),
        "catatan_terisi": sum(bool(grid[b["baris_masukan"] - 1][INPUT_COLS["catatan"] - 1].strip()) for b in blocks),
        "pasangan": pairs,
        "keputusan": [{"butir": int(d["butir"]), "jenis": d["jenis"], "keputusan": d["keputusan"], "ragu": d["ragu"]} for d in decisions],
        "l6_id": {n: b["l6_id"] for n, b in by_n.items()}})
    BLOCK_RESULT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k not in ("keputusan", "l6_id", "pasangan")}, ensure_ascii=False, indent=2))
    return 0


def key_comparison() -> int:
    """Keputusan butir baru vs kandidat teratas tiap pencarian. Kunci dibuka HANYA di sini, setelah hasil tersimpan."""
    if not BLOCK_RESULT_PATH.exists():
        raise SystemExit(f"{BLOCK_RESULT_PATH.name} belum ada: jalankan --read dulu (kunci dibuka hanya setelah hasil tersimpan)")
    res = json.loads(BLOCK_RESULT_PATH.read_text(encoding="utf-8"))
    key = json.loads(KEY_PATH.read_text(encoding="utf-8"))
    if {str(n): aid for n, aid in enumerate(key["urutan_butir"], 1)} != res["l6_id"]:
        raise SystemExit("urutan butir di kunci tidak sama dengan peta blok templat")
    decisions = [{**d, "butir": str(d["butir"])} for d in res["keputusan"]]
    report = {"catatan": "hanya untuk menilai pencarian otomatis; TIDAK dipakai untuk mengubah label",
              "butir_ragu": res["butir_ragu"], **compare_with_key(decisions, res["l6_id"], key["skor"])}
    KEY_RESULT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--make", action="store_true", help="buat berkas konfirmasi butir sisa (menolak menimpa)")
    g.add_argument("--read", action="store_true", help="baca keputusan dari sel blok berkas terisi dan hitung hasil")
    g.add_argument("--kunci", action="store_true", help="bandingkan keputusan dengan kandidat teratas (setelah --read)")
    ap.add_argument("--berkas", type=Path, default=FILLED_PATH, help=f"berkas terisi untuk --read (bawaan {FILLED_PATH.name})")
    args = ap.parse_args()
    return make() if args.make else key_comparison() if args.kunci else read(args.berkas)


if __name__ == "__main__":
    raise SystemExit(main())
