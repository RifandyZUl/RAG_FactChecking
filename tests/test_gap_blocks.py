"""Uji offline candidates.gap_blocks: tata letak blok, proteksi dan rumus Rekap, pembacaan dari sel blok."""

from openpyxl import load_workbook

from candidates.gap_blocks import (
    INPUT_COLS,
    NO_NARASI,
    SHEET,
    SUMMARY,
    combine,
    doubt_bounds,
    evidence_hash,
    layout,
    read_grid,
    read_inputs,
    write_workbook,
)


def _items(n: int = 3) -> list[dict]:
    return [{"id": f"L{i}", "judul": f"judul semu {i}", "tanggal": f"2026-01-0{i}", "pesan": "=pesan semu\nbaris dua",
             "ujung_rentang": i == 2,
             "kandidat": [{"id": str(100 + i), "tanggal": "02/01/2026", "label": "SALAH", "asal": "indeks", "judul": "t",
                           "narasi": "narasi semu", "kesimpulan": "k"},
                          {"id": str(200 + i), "tanggal": "05/03/2026", "label": "PENIPUAN", "asal": "antrean (belum di-ingest)",
                           "judul": "u", "narasi": "", "kesimpulan": "k2"}]} for i in range(1, n + 1)]


def _filled(tmp_path, inputs: dict[int, tuple[str, str]], tamper=None):
    """Tulis templat, isi sel blok seperti manusia (KEPUTUSAN, ragu), baca ulang sebagai kisi."""
    grid, blocks = layout(_items())
    path = tmp_path / "k.xlsx"
    write_workbook(grid, blocks, path)
    wb = load_workbook(path)
    for b in blocks:
        decision, doubt = inputs.get(b["butir"], ("", ""))
        wb[SHEET].cell(row=b["baris_masukan"], column=INPUT_COLS["KEPUTUSAN"], value=decision or None)
        wb[SHEET].cell(row=b["baris_masukan"], column=INPUT_COLS["ragu"], value=doubt or None)
    if tamper:
        tamper(wb)
    wb.save(path)
    return grid, blocks, read_grid(path)


def test_layout_orders_candidates_by_date_and_marks_missing_narasi():
    grid, blocks = layout(_items())
    assert [c["id"] for c in blocks[0]["kandidat"]] == ["201", "101"], "terbaru dulu, bukan menurut skor"
    assert blocks[0]["kandidat"][0]["tanpa_narasi"] and not blocks[0]["kandidat"][1]["tanpa_narasi"]
    assert sum(NO_NARASI in line for line in grid) == 3
    assert "skor" not in " ".join(" ".join(line) for line in grid).lower()
    assert all(grid[b["baris_masukan"] - 1][c - 1] == "" for b in blocks for c in INPUT_COLS.values()), "templat keluar kosong"


def test_workbook_protects_everything_except_block_input_cells(tmp_path):
    grid, blocks = layout(_items())
    path = tmp_path / "k.xlsx"
    write_workbook(grid, blocks, path)
    wb = load_workbook(path)
    ws, rk = wb[SHEET], wb[SUMMARY]
    assert ws.protection.sheet and rk.protection.sheet and not rk.protection.password
    row = blocks[0]["baris_masukan"]
    assert all(not ws.cell(row=row, column=c).protection.locked for c in INPUT_COLS.values())
    assert ws.cell(row=row, column=1).protection.locked and ws.cell(row=row + 2, column=7).protection.locked
    assert all(rk.cell(row=r, column=c).protection.locked for r in range(1, rk.max_row + 1) for c in range(1, 7))
    assert ws.cell(row=3, column=2).data_type == "s", "teks berawalan '=' tetap teks"
    formulas = [rk.cell(row=r, column=c).value for r in range(2, 5) for c in (2, 3)]
    assert all(isinstance(f, str) and f.startswith("=") and SHEET in f for f in formulas), "Rekap hanya rumus ke sel blok"
    assert read_grid(path) == grid[:-1] or evidence_hash(read_grid(path), blocks) == evidence_hash(grid, blocks)


def test_read_inputs_from_block_cells_normalizes_case_only(tmp_path):
    grid, blocks, back = _filled(tmp_path, {1: ("101", ""), 2: (" Tidak Ada ", "ya"), 3: ("bukan putusan hoaks", "")})
    assert evidence_hash(back, blocks) == evidence_hash(grid, blocks), "mengisi sel blok tidak mengubah bukti"
    res = read_inputs(back, blocks)
    assert res["perlu_ditanyakan"] == []
    assert [(d["jenis"], d["keputusan"], d["ragu"]) for d in res["butir"]] == [
        ("sama", "101", False), ("tidak_ada", "TIDAK ADA", True), ("bukan_putusan_hoaks", "BUKAN PUTUSAN HOAKS", False)]


def test_read_inputs_reports_instead_of_guessing(tmp_path):
    _, blocks, back = _filled(tmp_path, {1: ("999", ""), 2: ("TIDAK ADA", "mungkin"), 3: ("BUKAN PUTUSAN HOAKS", "YA")})
    problems = {(p["butir"], p["masalah"]) for p in read_inputs(back, blocks)["perlu_ditanyakan"]}
    assert problems == {("1", "id_di_luar_kandidat"), ("2", "ragu selain YA/kosong"), ("3", "ragu pada BUKAN PUTUSAN HOAKS")}
    _, blocks, empty = _filled(tmp_path, {1: ("101", "")})
    assert {p["masalah"] for p in read_inputs(empty, blocks)["perlu_ditanyakan"]} == {"kosong"}


def test_changed_evidence_is_detected(tmp_path):
    def tamper(wb):
        wb[SHEET].cell(row=7, column=7, value="narasi diubah")

    grid, blocks, back = _filled(tmp_path, {1: ("101", "")}, tamper)
    assert evidence_hash(back, blocks) != evidence_hash(grid, blocks)


def test_doubt_bounds_and_stopping_rule():
    def dec(n, kind, doubt=False):
        return {"butir": str(n), "jenis": kind, "ragu": doubt}

    new = doubt_bounds([dec(1, "sama", True), dec(2, "sama"), dec(3, "tidak_ada", True), dec(4, "tidak_ada"),
                        dec(5, "bukan_putusan_hoaks")])
    assert (new["utama"]["tidak_ada"], new["utama"]["butir_sah"]) == (2, 4)
    assert (new["batas_atas"]["tidak_ada"], new["batas_atas"]["butir_dibalik"]) == (3, [1])
    assert (new["batas_bawah"]["tidak_ada"], new["batas_bawah"]["butir_dibalik"]) == (1, [3])

    def part(k, n):
        return {"tidak_ada": k, "butir_sah": n}

    old = {"hasil_utama_bagian_11_2": part(29, 50),
           "sensitivitas_EKSPLORATIF_tidak_dipraregistrasi": {"batas_atas": part(37, 50), "batas_bawah": part(22, 50)}}
    amb = combine(old, {"utama": part(25, 49), "batas_atas": part(27, 49), "batas_bawah": part(24, 49)})
    assert (amb["hasil_utama"]["tidak_ada"], amb["hasil_utama"]["butir_sah"], amb["hasil_utama"]["status"]) == (54, 99, "ambigu")
    assert "TIDAK terpenuhi" in amb["keputusan_menurut_11_6"] and "tidak ada perluasan ketiga" in amb["keputusan_menurut_11_6"]
    assert amb["sensitivitas_EKSPLORATIF"]["batas_atas"]["tidak_ada"] == 64
    clear = combine(old, {"utama": part(40, 49), "batas_atas": part(40, 49), "batas_bawah": part(40, 49)})
    assert clear["hasil_utama"]["status"] == "jelas_di_atas_50" and clear["keputusan_menurut_11_6"] == "layak dikerjakan"
