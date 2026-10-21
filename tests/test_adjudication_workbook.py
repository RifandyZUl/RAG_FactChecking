"""Uji offline candidates.adjudication_workbook: pemilihan butir, opsi tanpa penanda, kunci terpisah, pembacaan."""

import hashlib
import json
import random

import pytest
from openpyxl import load_workbook

from candidates import adjudication_workbook as aw
from candidates.gap_blocks import INPUT_COLS, SHEET


def _final():
    def it(kode, art):
        return {"kode": kode, "l6_id": f"L{kode}", "dikeluarkan": None,
                "label_final": {"expected_verdict": "ditemukan" if art else "belum_ditemukan", "expected_article": art, "ragu": False}}
    return {"butir": [it("P1-1", "11"), it("P1-2", "12"), it("P1-3", None), it("P2-4", None), it("P2-5", "15"),
                      {**it("P1-4", None), "dikeluarkan": "K1"}]}


MODES = {"P1-1": {"verdict": "ditemukan", "article_id": "11"},      # sepakat -> tidak masuk
         "P1-2": {"verdict": "tidak_ditemukan", "article_id": None},  # penolakan palsu
         "P1-3": {"verdict": "ditemukan", "article_id": "13"},      # kecocokan palsu
         "P2-4": {"verdict": "tidak_ditemukan", "article_id": None},  # sepakat
         "P2-5": {"verdict": "ditemukan", "article_id": "16"},      # artikel salah
         "P1-4": {"verdict": "ditemukan", "article_id": "99"}}      # dikeluarkan K1


def test_disagreements_select_only_items_that_differ_from_final_label():
    rows = aw.disagreements(_final(), MODES)
    assert [(r["kode"], r["jenis"], r["opsi_label"], r["opsi_sistem"]) for r in rows] == [
        ("P1-2", "penolakan_palsu", "12", aw.NONE), ("P1-3", "kecocokan_palsu", aw.NONE, "13"),
        ("P2-5", "artikel_salah", "15", "16")]


@pytest.fixture
def adj(tmp_path, monkeypatch):
    for name, fname in [("ADJ_DIR", ""), ("XLSX_PATH", "a.xlsx"), ("FILLED_PATH", "a_terisi.xlsx"), ("TEMPLATE_PATH", "fp.json"),
                        ("KEY_PATH", "kunci.json"), ("RESULT_PATH", "hasil.json"), ("README_PATH", "BACA.txt"),
                        ("ARCHIVE_DIR", "arsip"), ("ARCHIVE_LOG", "arsip/log.jsonl")]:
        monkeypatch.setattr(aw, name, tmp_path / fname if fname else tmp_path)
    rows = aw.disagreements(_final(), MODES)
    valid = {r["l6_id"]: {"id": r["l6_id"], "judul": f"judul semu {i}", "tanggal": "2026-01-05", "pesan": "pesan semu"}
             for i, r in enumerate(rows)}
    by_id = {a: {"date": "01/02/2026", "label": "SALAH", "asal": "indeks", "title": f"t{a}", "narasi": "n", "kesimpulan": "k"}
             for a in ("12", "13", "15", "16")}
    fp = aw.write_adjudication(rows, valid, by_id, random.Random(3), {"x": "y"})
    monkeypatch.setattr(aw, "RECORDED_SHA256", {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                                                 for p in (aw.XLSX_PATH, aw.KEY_PATH, aw.TEMPLATE_PATH)})
    return tmp_path, fp


def _fill(tmp_path, fp, decide):
    wb = load_workbook(tmp_path / "a.xlsx")
    for b in fp["blok"]:
        wb[SHEET].cell(row=b["baris_masukan"], column=INPUT_COLS["KEPUTUSAN"], value=decide(b))
    out = tmp_path / "a_terisi.xlsx"
    wb.save(out)
    return out


def test_workbook_has_no_system_marker_and_template_has_no_mapping(adj):
    tmp_path, fp = adj
    wb = load_workbook(tmp_path / "a.xlsx")
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None).lower()
    assert not any(m in text for m in ("sistem", "label final", "p1-", "p2-", "penolakan palsu", "kecocokan palsu"))
    raw = (tmp_path / "fp.json").read_text(encoding="utf-8")
    assert "opsi_sistem" not in raw and "l6_id" not in raw and "P1-" not in raw
    assert sorted(len(b["kandidat"]) for b in fp["blok"]) == [1, 1, 2]


def test_read_maps_choice_to_label_or_system_side(adj):
    tmp_path, fp = adj
    key = {m["butir_adjudikasi"]: m for m in json.loads((tmp_path / "kunci.json").read_text(encoding="utf-8"))["peta"]}
    path = _fill(tmp_path, fp, lambda b: key[b["butir"]]["opsi_sistem"])
    assert aw.read(path) == 0
    res = json.loads((tmp_path / "hasil.json").read_text(encoding="utf-8"))
    assert res["memihak"] == {"sistem": 3}
    assert len([p for p in (tmp_path / "arsip").iterdir() if p.suffix == ".xlsx"]) == 1


def test_read_refuses_blank_and_unrecorded_hashes(adj, monkeypatch):
    tmp_path, fp = adj
    path = _fill(tmp_path, fp, lambda b: None if b["butir"] == 1 else aw.NONE)
    with pytest.raises(SystemExit, match="TIDAK dipetakan"):
        aw.read(path)
    assert not (tmp_path / "hasil.json").exists()
    monkeypatch.setattr(aw, "RECORDED_SHA256", {})
    v2 = tmp_path / "a_terisi_v2.xlsx"
    v2.write_bytes(path.read_bytes())
    with pytest.raises(SystemExit, match="hash tercatat"):
        aw.read(v2)


def test_recorded_hashes_match_design_document():
    from pathlib import Path
    doc = (Path(aw.__file__).parents[2] / "docs" / "rancangan_evaluasi_eksternal.md").read_text(encoding="utf-8")
    assert aw.RECORDED_SHA256 and all(f"`{n}`" in doc and h in doc for n, h in aw.RECORDED_SHA256.items())


def test_make_refuses_to_overwrite(adj):
    with pytest.raises(SystemExit):
        aw.make()


def _labels():
    def lab(art):
        return {"positif": art is not None, "sasaran": {art} if art else set(), "ragu": False, "putaran": 1, "strata_selisih": None}
    return {"P1-1": lab("11"), "P1-2": lab("12"), "P1-3": lab(None), "P2-4": lab(None), "P2-5": lab("15")}


DECISIONS = [  # P1-2 penolakan palsu -> TIDAK ADA (label keliru); P1-3 kecocokan palsu -> 13, ragu; P2-5 tetap label
    {"kode": "P1-2", "jenis_ketidaksesuaian": "penolakan_palsu", "keputusan": aw.NONE, "memihak": "sistem", "ragu": False},
    {"kode": "P1-3", "jenis_ketidaksesuaian": "kecocokan_palsu", "keputusan": "13", "memihak": "sistem", "ragu": True},
    {"kode": "P2-5", "jenis_ketidaksesuaian": "artikel_salah", "keputusan": "15", "memihak": "label", "ragu": False}]


def test_adjudicated_labels_apply_choice_except_label_side_and_doubtful_when_kept():
    full = aw.adjudicated_labels(_labels(), DECISIONS, keep_doubtful=False)
    assert not full["P1-2"]["positif"] and full["P1-3"]["sasaran"] == {"13"} and full["P2-5"]["sasaran"] == {"15"}
    kept = aw.adjudicated_labels(_labels(), DECISIONS, keep_doubtful=True)
    assert not kept["P1-2"]["positif"] and not kept["P1-3"]["positif"]
    assert _labels()["P1-2"]["positif"]  # label asli tidak diubah di tempat


def test_sensitivity_counts_both_directions_and_reports_worst_category():
    modes = {k: v for k, v in MODES.items() if k != "P1-4"}
    s = aw.sensitivity(_labels(), modes, {k: [] for k in modes}, DECISIONS)
    assert s["per_jenis"]["kecocokan_palsu"] == {"butir": 1, "kesalahan_label": ["P1-3"], "tetap_kesalahan_sistem": []}
    assert s["per_jenis"]["artikel_salah"]["tetap_kesalahan_sistem"] == ["P2-5"]
    assert (s["utama"]["positif"]["k"], s["utama"]["positif"]["n"]) == (1, 3)
    full, kept = s["disesuaikan"]["ragu_ikut_adjudikasi"], s["disesuaikan"]["ragu_label_tetap"]
    assert (full["positif"]["k"], full["positif"]["n"], full["kecocokan_palsu"]["k"], full["kecocokan_palsu"]["n"]) == (2, 3, 0, 2)
    assert (kept["positif"]["k"], kept["positif"]["n"], kept["kecocokan_palsu"]["k"], kept["kecocokan_palsu"]["n"]) == (1, 2, 1, 3)
    assert s["ragu"] == ["P1-3"] and "menguntungkan sistem" in s["sifat"]
