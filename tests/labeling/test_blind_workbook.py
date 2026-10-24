"""Uji offline labeling.sessions.blind_workbook: isi 37 butir, pengacakan tanpa penanda, kunci terpisah, pembacaan."""

import ast
import hashlib
import json
import random
from pathlib import Path

import pytest
from openpyxl import load_workbook

from labeling.gap.gap_blocks import INPUT_COLS, SHEET, layout
from labeling.sessions import blind_workbook as bw
from labeling.sessions.blind_session import SELECTION_PATH


def _source() -> dict[int, dict[int, dict]]:
    """Butir semu untuk dua putaran (teks semu; tanpa data artikel asli)."""
    def item(r: int, n: int) -> dict:
        return {"id": f"L{r}{n:02d}", "judul": f"judul semu {r}-{n}", "tanggal": "2026-01-05", "pesan": f"pesan semu {r}-{n}",
                "ujung_rentang": False,
                "kandidat": [{"id": str(1000 * r + 10 * n + k), "tanggal": f"0{k + 1}/02/2026", "label": "SALAH", "asal": "indeks",
                              "judul": f"t{k}", "narasi": f"narasi semu {k}", "kesimpulan": "k"} for k in range(5)]}
    return {r: {n: item(r, n) for n in range(1, 51)} for r in (1, 2)}


@pytest.fixture
def session(tmp_path, monkeypatch):
    for name, fname in [("SESSION_DIR", ""), ("XLSX_PATH", "s.xlsx"), ("FILLED_PATH", "s_terisi.xlsx"),
                        ("TEMPLATE_PATH", "fp.json"), ("KEY_PATH", "kunci.json"), ("RESULT_PATH", "hasil.json"),
                        ("README_PATH", "BACA.txt"), ("ARCHIVE_DIR", "arsip"), ("ARCHIVE_LOG", "arsip/log.jsonl")]:
        monkeypatch.setattr(bw, name, tmp_path / fname if fname else tmp_path)
    refs = bw.session_refs(json.loads(SELECTION_PATH.read_text(encoding="utf-8")))
    fp = bw.write_session(refs, _source(), random.Random(1), "x")
    monkeypatch.setattr(bw, "RECORDED_SHA256", {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                                                 for p in (bw.XLSX_PATH, bw.KEY_PATH, bw.TEMPLATE_PATH)})
    return tmp_path, refs, fp


def _fill(tmp_path: Path, fp: dict, decide) -> Path:
    wb = load_workbook(tmp_path / "s.xlsx")
    for b in fp["blok"]:
        wb[SHEET].cell(row=b["baris_masukan"], column=INPUT_COLS["KEPUTUSAN"], value=decide(b))
    out = tmp_path / "s_terisi.xlsx"
    wb.save(out)
    return out


def test_committed_selection_gives_30_retest_and_7_audit_items():
    refs = bw.session_refs(json.loads(SELECTION_PATH.read_text(encoding="utf-8")))
    assert len(refs) == 37 and len({r["kode"] for r in refs}) == 37
    assert sum(r["jenis"] == "uji_ulang" for r in refs) == 30
    assert sorted(r["kode"] for r in refs if r["jenis"] == "audit") == sorted(
        ["P1-10", "P1-20", "P1-29", "P1-30", "P1-38", "P1-45", "P2-2"])
    assert sum(r["putaran"] == 2 and r["jenis"] == "uji_ulang" for r in refs) == 20


def test_selection_that_differs_from_seeded_draw_is_refused():
    sel = json.loads(SELECTION_PATH.read_text(encoding="utf-8"))
    sel["uji_ulang"]["putaran_2"][0] = 2
    with pytest.raises(SystemExit):
        bw.session_refs(sel)


def test_layout_can_keep_given_candidate_order():
    item = _source()[1][1]
    _, blocks = layout([item], sort_by_date=False)
    assert [c["id"] for c in blocks[0]["kandidat"]] == [c["id"] for c in item["kandidat"]]
    _, blocks = layout([item])
    assert [c["id"] for c in blocks[0]["kandidat"]] != [c["id"] for c in item["kandidat"]], "bawaan tetap urut tanggal"


def test_session_file_has_no_marker_and_template_has_no_mapping(session):
    tmp_path, refs, fp = session
    wb = load_workbook(tmp_path / "s.xlsx")
    assert wb.sheetnames == ["konfirmasi", "Rekap", "petunjuk"]
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    assert not any(m in text.lower() for m in ("audit", "uji ulang", "uji_ulang", "p1-", "p2-", "putaran"))
    assert [b["butir"] for b in fp["blok"]] == list(range(1, 38))
    assert fp["masukan_terisi_saat_dibuat"] == 0
    raw_fp = (tmp_path / "fp.json").read_text(encoding="utf-8")
    assert "l6_id" not in raw_fp and "P1-" not in raw_fp and "audit" not in raw_fp
    key = json.loads((tmp_path / "kunci.json").read_text(encoding="utf-8"))["peta"]
    assert sorted(m["kode"] for m in key) == sorted(r["kode"] for r in refs)
    order = [m["kode"] for m in key]
    assert order != [r["kode"] for r in refs], "urutan butir diacak"
    src = _source()
    shuffled = sum([c["id"] for c in b["kandidat"]] != [c["id"] for c in src[m["putaran"]][m["butir_asli"]]["kandidat"]]
                   for b, m in zip(fp["blok"], key))
    assert shuffled > 0, "urutan kandidat diacak"
    for b, m in zip(fp["blok"], key):
        assert {c["id"] for c in b["kandidat"]} == {c["id"] for c in src[m["putaran"]][m["butir_asli"]]["kandidat"]}


def test_read_maps_decisions_to_original_items_only_when_complete(session):
    tmp_path, _, fp = session
    path = _fill(tmp_path, fp, lambda b: b["kandidat"][0]["id"] if b["butir"] % 2 else "tidak ada")
    assert bw.read(path) == 0
    res = json.loads((tmp_path / "hasil.json").read_text(encoding="utf-8"))
    key = {m["butir_sesi"]: m for m in json.loads((tmp_path / "kunci.json").read_text(encoding="utf-8"))["peta"]}
    assert res["jenis_keputusan"] == {"sama": 19, "tidak_ada": 18}
    assert res["butir_per_jenis"] == {"uji_ulang": 30, "audit": 7}
    for d in res["keputusan"]:
        assert d["kode"] == key[d["butir_sesi"]]["kode"]


def test_read_refuses_blank_items_and_changed_evidence(session):
    tmp_path, _, fp = session
    path = _fill(tmp_path, fp, lambda b: "TIDAK ADA" if b["butir"] != 5 else None)
    with pytest.raises(SystemExit):
        bw.read(path)
    assert not (tmp_path / "hasil.json").exists()
    wb = load_workbook(path)
    wb[SHEET].cell(row=fp["blok"][4]["baris_masukan"], column=INPUT_COLS["KEPUTUSAN"], value="TIDAK ADA")
    wb[SHEET].cell(row=2, column=2, value="judul diubah")
    wb.save(path)
    with pytest.raises(SystemExit):
        bw.read(path)
    assert not (tmp_path / "hasil.json").exists()


@pytest.mark.parametrize("name", ["s.xlsx", "kunci.json", "fp.json"])
def test_read_stops_when_a_recorded_hash_differs(session, name):
    tmp_path, _, fp = session
    path = _fill(tmp_path, fp, lambda b: "TIDAK ADA")
    target = tmp_path / name
    target.write_bytes(target.read_bytes() + b" ")
    with pytest.raises(SystemExit, match="hash tercatat"):
        bw.read(path)
    assert not (tmp_path / "hasil.json").exists()


def test_read_archives_the_file_before_any_validation(session):
    tmp_path, _, fp = session
    path = _fill(tmp_path, fp, lambda b: "TIDAK ADA" if b["butir"] != 5 else None)  # kosong -> validasi gagal
    want = hashlib.sha256(path.read_bytes()).hexdigest()
    with pytest.raises(SystemExit):
        bw.read(path)
    (copy,) = [p for p in (tmp_path / "arsip").iterdir() if p.suffix == ".xlsx"]
    assert hashlib.sha256(copy.read_bytes()).hexdigest() == want and want[:12] in copy.name
    log = [json.loads(x) for x in (tmp_path / "arsip" / "log.jsonl").read_text(encoding="utf-8").splitlines()]
    assert log == [{"waktu_utc": log[0]["waktu_utc"], "berkas": "s_terisi.xlsx", "sha256": want,
                    "ukuran": path.stat().st_size, "salinan": copy.name}]


def test_read_archives_even_when_template_hash_check_fails(session):
    tmp_path, _, fp = session
    path = _fill(tmp_path, fp, lambda b: "TIDAK ADA")
    (tmp_path / "kunci.json").write_bytes((tmp_path / "kunci.json").read_bytes() + b" ")
    with pytest.raises(SystemExit, match="hash tercatat"):
        bw.read(path)
    assert len([p for p in (tmp_path / "arsip").iterdir() if p.suffix == ".xlsx"]) == 1


def test_read_refuses_a_submitted_file_overwritten_under_the_same_name(session):
    tmp_path, _, fp = session
    path = _fill(tmp_path, fp, lambda b: "TIDAK ADA")
    assert bw.read(path) == 0
    assert bw.read(path) == 0, "membaca ulang berkas yang sama (hash sama) tetap boleh"
    (tmp_path / "hasil.json").unlink()
    wb = load_workbook(path)
    wb[SHEET].cell(row=fp["blok"][0]["baris_masukan"], column=INPUT_COLS["KEPUTUSAN"], value=fp["blok"][0]["kandidat"][0]["id"])
    wb.save(path)
    with pytest.raises(SystemExit, match="_v2"):
        bw.read(path)
    assert not (tmp_path / "hasil.json").exists()
    assert len([p for p in (tmp_path / "arsip").iterdir() if p.suffix == ".xlsx"]) == 3, "versi yang ditolak pun diarsipkan"
    v2 = tmp_path / "s_terisi_v2.xlsx"
    v2.write_bytes(path.read_bytes())
    assert bw.read(v2) == 0, "koreksi dengan nama baru dibaca"


def test_recorded_hashes_match_design_document():
    doc = (Path(__file__).resolve().parents[2] / "docs" / "rancangan" / "rancangan_evaluasi_eksternal.md").read_text(encoding="utf-8")
    assert all(f"`{name}`" in doc and h in doc for name, h in bw.RECORDED_SHA256.items())


def test_make_refuses_to_overwrite(session, monkeypatch):
    with pytest.raises(SystemExit):
        bw.make()


def test_module_code_does_not_touch_sealed_results_or_old_labels():
    tree = ast.parse(Path(bw.__file__).read_text(encoding="utf-8"))
    tree.body = tree.body[1:]  # docstring modul menyebut berkas yang TIDAK dibaca
    code = ast.unparse(tree)
    for name in ("external_eval", "eksternal_butir", "hasil_konfirmasi", "RESULT_PATH as", "kunci_kandidat", "_terisi.csv",
                 "gap_result", "KEY_PATH = OUT_DIR"):
        assert name not in code, name
    imported = {n.module for n in ast.walk(tree) if isinstance(n, ast.ImportFrom)}
    assert not any(m and m.startswith("evaluation") for m in imported)
