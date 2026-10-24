"""Uji offline labeling.sessions.v1_verification: isi tanpa penanda, jawaban SAMA/BEDA/RAGU, sensitivitas kemunduran."""

import hashlib
import json
import random
from pathlib import Path

import pytest
from openpyxl import load_workbook

from labeling.gap.gap_blocks import INPUT_COLS, SHEET
from labeling.sessions import v1_verification as vv

REG = {"non_batas": 44, "benar_ke_salah": ["v1-010", "v1-019", "v1-024", "v1-026", "v1-027"], "salah_ke_benar": [],
       "kemunduran_bersih": 5, "kategori": "mundur"}


def test_matched_articles_from_metrics():
    m = {"kemunduran_v1": {"perpindahan_non_batas": [{"id": i, "modus_1532": ["ditemukan", f"A{i}"]} for i in vv.ITEMS]}}
    assert vv.matched_articles(m) == {i: f"A{i}" for i in vv.ITEMS}
    m["kemunduran_v1"]["perpindahan_non_batas"][0]["modus_1532"] = ["tidak_ditemukan", None]
    with pytest.raises(SystemExit):
        vv.matched_articles(m)


@pytest.mark.parametrize("same,doubt,net,cat", [([], [], 5, "mundur"), (["v1-024"], [], 4, "mundur"),
                                                (["v1-024", "v1-026"], [], 3, "mundur"),
                                                (["v1-024", "v1-026", "v1-027"], [], 2, "tidak_mundur")])
def test_regression_sensitivity_excludes_same_and_keeps_threshold(same, doubt, net, cat):
    out = vv.regression_sensitivity(REG, same, doubt)
    assert out["SAMA_dikeluarkan"]["kemunduran_bersih"] == net and out["SAMA_dikeluarkan"]["kategori"] == cat
    assert out["utama_tidak_berubah"] == {"kemunduran_bersih": 5, "kategori": "mundur"}


def test_doubtful_is_reported_both_ways():
    out = vv.regression_sensitivity(REG, ["v1-024"], ["v1-026"])
    assert out["SAMA_dikeluarkan"]["kemunduran_bersih"] == 4 and out["RAGU_juga_dikeluarkan"]["kemunduran_bersih"] == 3


@pytest.fixture
def ver(tmp_path, monkeypatch):
    for name, fname in [("VER_DIR", ""), ("XLSX_PATH", "v.xlsx"), ("FILLED_PATH", "v_terisi.xlsx"), ("TEMPLATE_PATH", "fp.json"),
                        ("KEY_PATH", "kunci.json"), ("RESULT_PATH", "hasil.json"), ("README_PATH", "BACA.txt"),
                        ("ARCHIVE_DIR", "arsip"), ("ARCHIVE_LOG", "arsip/log.jsonl")]:
        monkeypatch.setattr(vv, name, tmp_path / fname if fname else tmp_path)
    items = [{"v1_id": i, "klaim": f"klaim semu {n}", "article_id": f"9{n}"} for n, i in enumerate(vv.ITEMS)]
    by_id = {f"9{n}": {"date": "01/02/2026", "label": "SALAH", "asal": "indeks", "title": f"t{n}", "narasi": "n", "kesimpulan": "k"}
             for n in range(3)}
    fp = vv.write_files(items, by_id, random.Random(2), {"x": "y"})
    monkeypatch.setattr(vv, "RECORDED_SHA256", {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                                                 for p in (vv.XLSX_PATH, vv.KEY_PATH, vv.TEMPLATE_PATH)})
    metrics = tmp_path / "m.json"
    metrics.write_text(json.dumps({"kemunduran_v1": REG}), encoding="utf-8")
    monkeypatch.setattr("evaluation.external_analysis.METRICS_PATH", metrics)
    return tmp_path, fp


def _fill(tmp_path: Path, fp, decide):
    wb = load_workbook(tmp_path / "v.xlsx")
    for b in fp["blok"]:
        wb[SHEET].cell(row=b["baris_masukan"], column=INPUT_COLS["KEPUTUSAN"], value=decide(b))
    out = tmp_path / "v_terisi.xlsx"
    wb.save(out)
    return out


def test_workbook_hides_item_ids_and_generator_output(ver):
    tmp_path, _ = ver
    wb = load_workbook(tmp_path / "v.xlsx")
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None).lower()
    assert not any(m in text for m in ("v1-0", "liputan6 terbit", "verdict", "ditemukan", "generator", "sistem"))
    assert "kueri_set_uji_v1" in text and "klaim semu 0" in text
    assert "v1-0" not in (tmp_path / "fp.json").read_text(encoding="utf-8")


def test_read_maps_answers_and_recomputes_regression(ver):
    tmp_path, fp = ver
    path = _fill(tmp_path, fp, lambda b: "sama" if b["butir"] == 1 else "BEDA")
    assert vv.read(path) == 0
    res = json.loads((tmp_path / "hasil.json").read_text(encoding="utf-8"))
    assert res["jawaban"] == {"SAMA": 1, "BEDA": 2}
    assert res["sensitivitas_kemunduran"]["SAMA_dikeluarkan"]["kemunduran_bersih"] == 4


def test_read_refuses_invalid_answer_and_unrecorded_hash(ver, monkeypatch):
    tmp_path, fp = ver
    path = _fill(tmp_path, fp, lambda b: "MUNGKIN")
    with pytest.raises(SystemExit, match="TIDAK dipetakan"):
        vv.read(path)
    monkeypatch.setattr(vv, "RECORDED_SHA256", {})
    v2 = tmp_path / "v_terisi_v2.xlsx"
    v2.write_bytes(path.read_bytes())
    with pytest.raises(SystemExit, match="hash tercatat"):
        vv.read(v2)


def test_recorded_hashes_match_design_document():
    doc = (Path(__file__).resolve().parents[2] / "docs" / "rancangan" / "rancangan_evaluasi_eksternal.md").read_text(encoding="utf-8")
    assert vv.RECORDED_SHA256 and all(f"`{n}`" in doc and h in doc for n, h in vv.RECORDED_SHA256.items())
