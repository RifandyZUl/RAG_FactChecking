"""Uji lembar uji coba pedoman 2.0 (labeling.sessions.guide_trial): pemilihan berbenih, isi lembar, dan sel masukan."""

import json
import random
from collections import Counter
from pathlib import Path

import pytest
from openpyxl import load_workbook

from labeling.gap.gap_workbook import provenance
from labeling.sessions.guide_trial import ALLOCATION, SHEET, build, select, stratum


def fake_strata() -> dict[str, tuple[str, str]]:
    kinds = [("positif", "penipuan")] * 8 + [("negatif", "penipuan")] * 8 + [("positif", "non_penipuan")] * 4 \
        + [("negatif", "non_penipuan")] * 4
    return {f"P{1 + i % 2}-{i + 1}": s for i, s in enumerate(kinds)}


def test_stratum_uses_final_verdict_and_candidate_labels() -> None:
    assert stratum({"expected_verdict": "ditemukan"}, ["SALAH", "PENIPUAN"]) == ("positif", "penipuan")
    assert stratum({"expected_verdict": "belum_ditemukan"}, ["SALAH", "PARODI"]) == ("negatif", "non_penipuan")


def test_select_is_seeded_and_follows_allocation() -> None:
    strata = fake_strata()
    first, again = select(strata), select(dict(reversed(list(strata.items()))))
    assert first == again  # urutan masukan tidak berpengaruh
    assert Counter(strata[c] for c in first) == Counter(ALLOCATION)
    assert len(set(first)) == sum(ALLOCATION.values())
    assert select(strata, seed=1) != first


def test_select_refuses_small_stratum() -> None:
    strata = {k: v for k, v in fake_strata().items() if v != ("negatif", "non_penipuan")}
    with pytest.raises(SystemExit):
        select(strata)


def item(n: int) -> dict:
    return {"judul": f"Cek Fakta: contoh {n}", "tanggal": "2026-09-20", "pesan": "=pesan contoh berawalan sama-dengan",
            "kandidat": [{"id": f"{n}00", "tanggal": "01/09/2026", "label": "PENIPUAN", "judul": "[PENIPUAN] Judul",
                          "narasi": "narasi", "kesimpulan": "kesimpulan", "referensi": ""},
                         {"id": f"{n}01", "tanggal": "02/09/2026", "label": "SALAH", "judul": "[SALAH] Judul",
                          "narasi": "", "kesimpulan": "kesimpulan", "referensi": "https://contoh.example/a"}]}


def test_build_writes_sheet_key_and_fingerprint_without_labels(tmp_path: Path) -> None:
    items = {"P1-1": item(1), "P2-7": item(2)}
    fp = build(items, ["P1-1", "P2-7"], random.Random(3), out_dir=tmp_path)
    xlsx = tmp_path / "uji_coba_pedoman_v2.xlsx"
    ws = load_workbook(xlsx)[SHEET]
    text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    for leaked in ("P1-1", "P2-7", "ditemukan", "positif", "negatif", "label_final"):
        assert leaked not in text
    assert "(TIDAK ADA DI SUMBER)" in text  # Narasi dan referensi kosong ditandai
    assert ws.cell(row=5, column=2).data_type == "s"  # pesan berawalan "=" tetap teks
    # sel gabungan (MergedCell) mengikuti sel kiri-atasnya, jadi tidak dihitung
    open_cells = {(c.row, c.column) for row in ws.iter_rows() for c in row
                  if not c.protection.locked and type(c).__name__ != "MergedCell"}
    assert open_cells == {tuple(x) for x in fp["sel_masukan"]}
    assert ws.protection.sheet
    assert all(not ws.cell(row=r, column=c).value for r, c in open_cells)  # sel masukan kosong
    assert provenance(xlsx)["creator"].startswith("labeling.sessions.guide_trial")
    key = json.loads((tmp_path / "kunci_uji_coba.json").read_text(encoding="utf-8"))
    assert sorted(m["kode"] for m in key["peta"]) == ["P1-1", "P2-7"]
    assert "kode" not in json.dumps(fp["blok"])  # sidik jari templat tanpa peta ke butir asli
