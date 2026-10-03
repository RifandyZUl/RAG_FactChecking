"""Uji berkas konfirmasi .xlsx (candidates.gap_workbook): tata letak, asal-usul, dan pembacaan keputusan. Tanpa model."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from candidates.gap_review import COLUMNS
from candidates.gap_workbook import (
    SHEET,
    content_hash,
    normalize_decision,
    proportion,
    provenance,
    read_decisions,
    read_workbook_rows,
    summarize,
    write_workbook,
)


def row(butir: int, kandidat: int, tbh_id: str, first: bool, **over: str) -> dict[str, str]:
    r = {"butir": str(butir), "L6_tanggal": "2026-09-20" if first else "", "ujung_rentang": "YA" if first and butir == 1 else "",
         "L6_judul": f"Cek Fakta: Contoh {butir}" if first else "", "L6_kutipan_pesan": "pesan contoh" if first else "",
         "KEPUTUSAN": "", "catatan": "", "kandidat": str(kandidat), "TBH_id": tbh_id, "TBH_tanggal": "05/01/2026",
         "TBH_label": "SALAH" if kandidat == 1 else "PENIPUAN", "TBH_asal": "indeks", "TBH_judul": f"[SALAH] Judul {tbh_id}",
         "TBH_kesimpulan": "Faktanya contoh. " * 30}
    return {**r, **over}


def sample_rows() -> list[dict[str, str]]:
    return [row(1, 1, "100", True), row(1, 2, "101", False), row(2, 1, "200", True), row(2, 2, "201", False),
            row(3, 1, "300", True)]


def test_workbook_layout_and_empty_decisions(tmp_path: Path) -> None:
    path = tmp_path / "k.xlsx"
    rows = sample_rows()
    rows[0]["L6_kutipan_pesan"] = "=SUM(1,2) pesan berawalan sama-dengan\x0b"
    stats = write_workbook(rows, path)
    ws = load_workbook(path)[SHEET]
    assert [c.value for c in ws[1]] == COLUMNS and ws.cell(1, ws.max_column).value == "TBH_kesimpulan", "Kesimpulan paling kanan"
    assert ws.freeze_panes == "A2", "baris judul dibekukan"
    assert all(c.alignment.wrap_text for r in ws.iter_rows() for c in r), "teks dibungkus"
    assert ws.column_dimensions["N"].width >= 60 and ws.column_dimensions["A"].width < 10
    assert ws.cell(2, 5).data_type == "s" and ws.cell(2, 5).value.startswith("=SUM"), "teks tidak boleh menjadi rumus"
    assert stats == {"karakter_kendali_dibuang": 1, "sel_dipotong": 0}
    back = read_workbook_rows(path)
    assert all(r["KEPUTUSAN"] == "" and r["catatan"] == "" for r in back), "keputusan tidak pernah diisi kode"
    assert [r["TBH_id"] for r in back] == ["100", "101", "200", "201", "300"]
    assert not any("skor" in c.lower() for c in COLUMNS)


def test_workbook_refuses_prefilled_decisions(tmp_path: Path) -> None:
    rows = sample_rows()
    rows[2]["KEPUTUSAN"] = "200"
    with pytest.raises(ValueError, match="sudah terisi"):
        write_workbook(rows, tmp_path / "k.xlsx")


def test_provenance_marks_code_generated_file(tmp_path: Path) -> None:
    path = tmp_path / "k.xlsx"
    write_workbook(sample_rows(), path)
    prov = provenance(path)
    assert "openpyxl" in prov["creator"] and "dihasilkan kode" in prov["creator"]
    assert prov["lastModifiedBy"] == "", "templat belum pernah disimpan aplikasi lain"


def test_content_hash_ignores_human_columns_only(tmp_path: Path) -> None:
    path = tmp_path / "k.xlsx"
    rows = sample_rows()
    write_workbook(rows, path)
    wb = load_workbook(path)
    wb[SHEET].cell(2, COLUMNS.index("KEPUTUSAN") + 1).value = 100  # Excel menyimpan id yang diketik sebagai angka
    wb[SHEET].cell(2, COLUMNS.index("catatan") + 1).value = "catatan manusia"
    wb.save(path)
    filled = read_workbook_rows(path)
    assert filled[0]["KEPUTUSAN"] == "100", "angka bulat dibaca sebagai id, bukan 100.0"
    assert content_hash(filled) == content_hash(rows)
    filled[1]["TBH_judul"] = "diubah"
    assert content_hash(filled) != content_hash(rows)


def test_decisions_are_read_not_guessed() -> None:
    rows = sample_rows() + [row(4, 1, "400", True), row(5, 1, "500", True), row(5, 2, "501", False), row(6, 1, "600", True)]
    for i, value in {0: " 101 ", 2: "tidak  ada", 4: "Bukan putusan hoaks.", 5: "999", 6: "", 8: "mungkin"}.items():
        rows[i]["KEPUTUSAN"] = value
    rows[7]["KEPUTUSAN"] = "501"  # diketik di baris kandidat, bukan baris pertama butir
    res = read_decisions(rows)
    kinds = {d["butir"]: d["jenis"] for d in res["butir"]}
    assert kinds == {"1": "sama", "2": "tidak_ada", "3": "bukan_putusan_hoaks", "4": "id_di_luar_kandidat", "5": "kosong",
                     "6": "tidak_terbaca"}
    assert res["butir"][0]["label_tbh"] == "PENIPUAN" and res["butir"][0]["ujung_rentang"]
    problems = {(p["butir"], p["masalah"]) for p in res["perlu_ditanyakan"]}
    assert problems == {("4", "id_di_luar_kandidat"), ("5", "kosong"), ("5", "keputusan di baris bukan-pertama"),
                        ("6", "tidak_terbaca")}
    assert normalize_decision("  tidak\n ada ") == "TIDAK ADA"


def test_proportion_status_follows_wilson_rule() -> None:
    assert proportion(40, 49)["status"] == "jelas_di_atas_50"
    assert proportion(10, 49)["status"] == "jelas_di_bawah_50"
    assert proportion(28, 49)["status"] == "ambigu", "proporsi teramati di atas 50% tetapi interval melintasi 50%"
    assert proportion(0, 0)["status"] == "tanpa_butir"


def test_summary_excludes_non_hoax_items_from_denominator() -> None:
    decisions = [{"butir": "1", "jenis": "tidak_ada", "keputusan": "TIDAK ADA", "ujung_rentang": True, "label_tbh": ""},
                 {"butir": "2", "jenis": "sama", "keputusan": "200", "ujung_rentang": False, "label_tbh": "SALAH"},
                 {"butir": "3", "jenis": "bukan_putusan_hoaks", "keputusan": "", "ujung_rentang": False, "label_tbh": ""},
                 {"butir": "4", "jenis": "tidak_ada", "keputusan": "TIDAK ADA", "ujung_rentang": False, "label_tbh": ""}]
    s = summarize(decisions)
    assert (s["semua"]["tidak_ada"], s["semua"]["butir_sah"]) == (2, 3)
    assert (s["ujung_rentang"]["tidak_ada"], s["ujung_rentang"]["butir_sah"]) == (1, 1)
    assert (s["di_luar_ujung_rentang"]["tidak_ada"], s["di_luar_ujung_rentang"]["butir_sah"]) == (1, 2)
    assert s["label_tbh_pada_pasangan_sama"] == {"SALAH": 1}
