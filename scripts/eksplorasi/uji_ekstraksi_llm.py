"""
Uji ketepatan ekstraksi unsur inti dengan LLM (EKSPLORATIF; rancangan Versi 2 bagian A2.6).

Yang diukur adalah KETEPATAN EKSTRAKSI per unsur (diperiksa manual oleh pemilik proyek), bukan jumlah kecocokan
palsu yang tertangkap. Rencana, prompt, format keluaran, cara pemeriksaan, dan kriteria lanjut/berhenti ditulis
dan di-commit SEBELUM modul ini dijalankan (A2.6).

Sumber (lokal, tidak di-commit): 47 pesan (kutipan Liputan6 dari butir evaluasi eksternal yang dijawab
"ditemukan") dan 45 artikel TurnBackHoax yang dipilih sistem untuk butir itu (judul + Narasi + Kesimpulan).
Teks dan hasil ekstraksi hanya disimpan di `data/eksplorasi/` (di-gitignore). `src/generator.py` tidak disentuh;
modul ini hanya memakai lapisan penyedia `src/llm/` (pembatas laju dan buku besar kuota yang sama).

Pemakaian (dari root proyek):
  python scripts/eksplorasi/uji_ekstraksi_llm.py --anggaran          # sisa kuota, tanpa panggilan
  python scripts/eksplorasi/uji_ekstraksi_llm.py --jalankan --run 1  # 92 panggilan; bisa dilanjutkan
  python scripts/eksplorasi/uji_ekstraksi_llm.py --jalankan --run 2  # 92 panggilan (kestabilan antar-run)
  python scripts/eksplorasi/uji_ekstraksi_llm.py --lembar            # lembar pemeriksaan manual dari run 1
  python scripts/eksplorasi/uji_ekstraksi_llm.py --kestabilan        # run 1 vs run 2, tanpa panggilan
  python scripts/eksplorasi/uji_ekstraksi_llm.py --hitung BERKAS     # setelah lembar diisi pemilik proyek
"""

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

OUT_DIR = ROOT / "data" / "eksplorasi"
SHEET_PATH = OUT_DIR / "lembar_pemeriksaan_ekstraksi.xlsx"

UNSUR = ("subjek", "tawaran", "tindakan", "nominal", "saluran", "periode")

SYSTEM_PROMPT = """Anda mengekstrak unsur inti dari satu teks berbahasa Indonesia: sebuah pesan yang beredar, atau sebuah
artikel cek fakta tentang pesan yang beredar. Tugas Anda HANYA mencatat apa yang tertulis eksplisit di teks.
Jangan menilai benar atau salah, jangan menebak, dan jangan menambahkan pengetahuan dari luar teks.

Untuk artikel cek fakta: ambil unsur dari KLAIM YANG DIPERIKSA (isi pesan atau unggahan yang dibahas), bukan dari
penjelasan pemeriksa fakta, bantahan, atau nama lembaga yang dikutip sebagai sumber bantahan.

Unsur yang diambil:
- subjek: lembaga, instansi, perusahaan, atau tokoh yang disebut sebagai pihak yang menawarkan, mengumumkan,
  atau melakukan sesuatu dalam klaim. Tulis namanya seperti di teks dan beri jenis "lembaga" atau "tokoh".
  Jangan masukkan pihak yang hanya disebut sebagai latar.
- tawaran: program, bantuan, lowongan, hadiah, atau keuntungan yang ditawarkan atau diklaim. Bila klaimnya bukan
  tawaran (mis. kabar peristiwa), tulis pokok klaimnya dalam frasa pendek. Kosong bila tidak ada.
- tindakan: apa yang diklaim terjadi atau diminta dari pembaca, dalam frasa pendek (mis. "pendaftaran dibuka",
  "membagikan dana", "diminta mengisi data"). Kosong bila tidak ada.
- nominal: semua jumlah uang yang disebut, ditulis seperti di teks. Jumlah suka, komentar, atau tayangan BUKAN
  nominal. Kosong bila tidak ada.
- saluran: tautan, domain, nomor telepon atau WhatsApp, atau nama aplikasi tempat pembaca diminta mendaftar atau
  menghubungi, ditulis seperti di teks. Kosong bila tidak ada.
- periode: tanggal, bulan, tahun, atau gelombang yang disebut sebagai waktu tawaran atau peristiwa (mis. "2026",
  "gelombang 3", "Maret"). Kosong bila tidak ada.

Bila sebuah unsur tidak disebut, kembalikan daftar kosong atau string kosong. Jangan mengarang."""

USER_TEMPLATE = "Jenis teks: {jenis}\n\nTeks:\n<<<\n{teks}\n>>>"

SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {
        "subjek": {"type": "array", "items": {"type": "object", "properties": {
            "nama": {"type": "string"}, "jenis": {"type": "string", "enum": ["lembaga", "tokoh"]}},
            "required": ["nama", "jenis"]}},
        "tawaran": {"type": "string"},
        "tindakan": {"type": "string"},
        "nominal": {"type": "array", "items": {"type": "string"}},
        "saluran": {"type": "array", "items": {"type": "string"}},
        "periode": {"type": "array", "items": {"type": "string"}},
    },
    "required": list(UNSUR),
}


def load_sources() -> list[dict[str, Any]]:
    """Sumber unik berurutan per butir: pesan butir, lalu artikelnya (artikel yang sudah muncul tidak diulang)."""
    from evaluation.external_eval import build_liputan6_cases, norm_sha

    frozen = {b["kode"]: b for b in json.loads((ROOT / "testset" / "eksternal_butir.json").read_text(encoding="utf-8"))["butir"]}
    results = json.loads((ROOT / "testset" / "eksternal_hasil_per_butir.json").read_text(encoding="utf-8"))["butir"]
    articles = {a["article_id"]: a for a in json.loads((ROOT / "data" / "articles.json").read_text(encoding="utf-8"))}
    cases = {c["kode"]: c for c in build_liputan6_cases()}
    out, seen = [], set()
    for r in results:
        if r["keputusan_modus"] != "ditemukan":
            continue
        kode = r["kode"]
        if norm_sha(cases[kode]["kutipan"]) != frozen[kode]["sha256_kutipan"]:
            raise SystemExit(f"{kode}: hash kutipan tidak cocok dengan testset/eksternal_butir.json")
        out.append({"sid": f"pesan:{kode}", "kode": kode, "sisi": "pesan", "teks": cases[kode]["kutipan"]})
        aid = r["article_id_modus"]
        if aid not in seen:
            seen.add(aid)
            a = articles[aid]
            teks = f"{a['title']}\n\nNarasi:\n{a['narasi']}\n\nKesimpulan:\n{a['kesimpulan']}"
            out.append({"sid": f"artikel:{aid}", "kode": kode, "sisi": "artikel", "teks": teks})
    return out


def run_path(run: int) -> Path:
    return OUT_DIR / f"ekstraksi_run{run}.jsonl"


def load_run(run: int) -> dict[str, dict]:
    p = run_path(run)
    if not p.exists():
        return {}
    return {r["sid"]: r for r in (json.loads(x) for x in p.read_text(encoding="utf-8").splitlines() if x.strip())}


def run_extraction(run: int, budget_only: bool) -> int:
    from llm import LLMError, LLMQuotaExhaustedError, get_provider
    from llm.ledger import plan_budget

    sources = load_sources()
    done = load_run(run)
    todo = [s for s in sources if s["sid"] not in done or not done[s["sid"]].get("ok")]
    provider = get_provider()
    print(f"run {run} | sumber {len(sources)} | selesai {len(sources) - len(todo)} | akan dijalankan {len(todo)} "
          f"| model {provider.model} thinking {getattr(provider, 'thinking_level', '?')}")
    if provider.ledger is not None and provider.limits is not None:
        plan = plan_budget(provider.ledger, provider.limits, len(todo), 4)
        print(plan.describe())
        if todo and not plan.ok:
            print("TIDAK DIMULAI: kuota harian tidak cukup.")
            return 4
    if budget_only or not todo:
        return 0
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    calls_before = len(provider.records)
    with run_path(run).open("a", encoding="utf-8", newline="\n") as fh:
        for i, s in enumerate(todo, 1):
            user = USER_TEMPLATE.format(jenis="pesan yang beredar" if s["sisi"] == "pesan" else "artikel cek fakta",
                                        teks=s["teks"])
            rec = {"sid": s["sid"], "kode": s["kode"], "sisi": s["sisi"], "run": run}
            try:
                raw = provider.generate(SYSTEM_PROMPT, user, json_schema=SCHEMA)
                data = json.loads(raw)
                rec |= {"ok": True, "unsur": {k: data.get(k, [] if k not in ("tawaran", "tindakan") else "") for k in UNSUR}}
            except LLMQuotaExhaustedError as e:  # kuota harian habis: berhenti, jangan kirim permintaan lagi
                print(f"BERHENTI: kuota harian habis ({e}); jalankan ulang perintah yang sama setelah reset")
                return 4
            except (LLMError, json.JSONDecodeError, AttributeError) as e:  # dicatat per sumber; dicoba lagi pada jalan berikutnya
                rec |= {"ok": False, "galat": f"{type(e).__name__}: {e}"[:300]}
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
            fh.flush()
            print(f"  [{i}/{len(todo)}] {s['sid']}: {'ok' if rec['ok'] else rec['galat']}")
    print(f"panggilan LLM pada jalan ini: {len(provider.records) - calls_before}")
    return 0


def _norm_value(v: Any) -> Any:
    if isinstance(v, str):
        return re.sub(r"\s+", " ", v.strip().lower())
    if isinstance(v, list):
        items = [f"{x.get('nama', '')}|{x.get('jenis', '')}" if isinstance(x, dict) else str(x) for x in v]
        return sorted(_norm_value(x) for x in items)
    return v


def _show(v: Any) -> str:
    if isinstance(v, list):
        if not v:
            return "(tidak disebut)"
        return "; ".join(f"{x['nama']} [{x['jenis']}]" if isinstance(x, dict) else str(x) for x in v)
    return v if v else "(tidak disebut)"


def make_sheet() -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    if SHEET_PATH.exists():
        raise SystemExit(f"{SHEET_PATH} sudah ada: tidak ditimpa")
    run1 = load_run(1)
    sources = load_sources()
    missing = [s["sid"] for s in sources if not run1.get(s["sid"], {}).get("ok")]
    if missing:
        raise SystemExit(f"run 1 belum lengkap: {len(missing)} sumber")
    wb = Workbook()
    ws = wb.active
    ws.title = "periksa"
    head = ["no", "butir", "sisi", "id_sumber", "teks_sumber"]
    for u in UNSUR:
        head += [f"{u}_hasil", f"{u}_benar (B/S)"]
    head += ["catatan"]
    ws.append(head)
    dv = DataValidation(type="list", formula1='"B,S"', allow_blank=True)
    ws.add_data_validation(dv)
    for n, s in enumerate(sources, 1):
        u = run1[s["sid"]]["unsur"]
        row = [n, s["kode"], s["sisi"], s["sid"], s["teks"]]
        for k in UNSUR:
            row += [_show(u[k]), ""]
        row += [""]
        ws.append(row)
        for j in range(len(UNSUR)):
            dv.add(ws.cell(row=n + 1, column=7 + 2 * j))
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.column_dimensions["E"].width = 70
    for col in "FHJLNP":
        ws.column_dimensions[col].width = 28
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "F2"
    pet = wb.create_sheet("petunjuk")
    for line in [
        "Isi kolom <unsur>_benar dengan B (benar) atau S (salah) untuk setiap unsur, dengan membandingkan hasil dengan teks sumber.",
        "B: semua yang tertulis eksplisit untuk unsur itu terambil, dan tidak ada tambahan yang tidak ada di teks.",
        "B juga untuk '(tidak disebut)' bila unsur itu memang tidak ada di teks.",
        "S: ada yang terlewat, ada tambahan yang tidak ada di teks, salah jenis (lembaga/tokoh), atau mengambil unsur dari bantahan pemeriksa fakta, bukan dari klaim.",
        "Kata-kata tidak harus persis; yang dinilai adalah isinya.",
        "Jangan mengubah kolom lain. Isi tanpa bantuan AI (aturan wajib #5); bila ragu, isi S dan tulis alasannya di catatan.",
    ]:
        pet.append([line])
    pet.column_dimensions["A"].width = 140
    wb.save(SHEET_PATH)
    print(f"lembar dibuat: {SHEET_PATH} ({len(sources)} baris: "
          f"{sum(s['sisi'] == 'pesan' for s in sources)} pesan, {sum(s['sisi'] == 'artikel' for s in sources)} artikel)")
    return 0


def stability() -> dict[str, Any]:
    r1, r2 = load_run(1), load_run(2)
    common = [sid for sid in r1 if r1[sid].get("ok") and r2.get(sid, {}).get("ok")]
    out = {"sumber_dibandingkan": len(common)}
    for sisi in ("pesan", "artikel"):
        ids = [sid for sid in common if r1[sid]["sisi"] == sisi]
        out[sisi] = {k: f"{sum(_norm_value(r1[i]['unsur'][k]) == _norm_value(r2[i]['unsur'][k]) for i in ids)}/{len(ids)}"
                     for k in UNSUR}
    return out


def compute(sheet: Path) -> int:
    from openpyxl import load_workbook

    ws = load_workbook(sheet)["periksa"]
    head = [c.value for c in ws[1]]
    rows = [dict(zip(head, (c.value for c in r))) for r in ws.iter_rows(min_row=2)]
    problems = [f"baris {i + 2} {r['id_sumber']} {u}" for i, r in enumerate(rows) for u in UNSUR
                if str(r.get(f"{u}_benar (B/S)") or "").strip().upper() not in ("B", "S")]
    if problems:
        print(f"BELUM DIHITUNG: {len(problems)} sel B/S kosong atau tidak sah, mis. {problems[:5]}")
        return 1
    acc = {}
    for sisi in ("pesan", "artikel"):
        sub = [r for r in rows if r["sisi"] == sisi]
        acc[sisi] = {u: (sum(str(r[f"{u}_benar (B/S)"]).strip().upper() == "B" for r in sub), len(sub)) for u in UNSUR}
    print(json.dumps({"ketepatan": {s: {u: f"{b}/{n}" for u, (b, n) in d.items()} for s, d in acc.items()},
                      "kestabilan": stability()}, ensure_ascii=False, indent=1))
    (OUT_DIR / "hasil_ekstraksi.json").write_text(json.dumps({"ketepatan": acc, "kestabilan": stability()},
                                                             ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    return 0


def main() -> int:
    from hf_env import prepare_hf_env

    prepare_hf_env()
    ap = argparse.ArgumentParser(description=__doc__.strip().splitlines()[0])
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--anggaran", action="store_true")
    g.add_argument("--jalankan", action="store_true")
    g.add_argument("--lembar", action="store_true")
    g.add_argument("--kestabilan", action="store_true")
    g.add_argument("--hitung", type=Path)
    ap.add_argument("--run", type=int, choices=(1, 2), default=1)
    a = ap.parse_args()
    if a.anggaran or a.jalankan:
        return run_extraction(a.run, budget_only=a.anggaran)
    if a.lembar:
        return make_sheet()
    if a.kestabilan:
        print(json.dumps(stability(), ensure_ascii=False, indent=1))
        return 0
    return compute(a.hitung)


if __name__ == "__main__":
    sys.exit(main())
